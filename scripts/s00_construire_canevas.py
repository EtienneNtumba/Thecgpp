"""
s00_construire_canevas.py - Construit le canevas Excel de la base de données CGPP (modèle vide).

Le canevas contient :
  - LISEZMOI : objet, légende des couleurs, circuit des données, liste des feuilles ;
  - dictionnaire : une ligne par colonne (type, description, obligatoire, valeurs autorisées) ;
  - listes : valeurs autorisées utilisées par les listes déroulantes ;
  - une feuille par table du modèle de données (en-têtes, formats, listes déroulantes,
    une ligne d'exemple sur les feuilles saisies à la main) ;
  - tableau_suivi : suivi des indicateurs calculé par formules (SOMME.SI.ENS sur saisie_hebdo),
    pourcentages recalculés à partir des numérateurs et dénominateurs.

Usage : python scripts/s00_construire_canevas.py
"""
from __future__ import annotations

import sys
from datetime import date, datetime
from pathlib import Path

import openpyxl
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, str(Path(__file__).resolve().parent))
from commun import LISTES, SCHEMA, charger_config, cible_numerique, colonnes, journal  # noqa: E402
from referentiel_indicateurs import matrice  # noqa: E402

LOG = journal("canevas")
POLICE = "Arial"
BLEU_FONCE, BLEU_PALE, JAUNE_SAISIE, GRIS_CALCUL = "1F3864", "DCE6F2", "FFF2CC", "EDEDED"
F_TITRE = Font(name=POLICE, size=14, bold=True, color=BLEU_FONCE)
F_SOUS = Font(name=POLICE, size=11, bold=True, color=BLEU_FONCE)
F_TXT = Font(name=POLICE, size=10)
F_ENTETE = Font(name=POLICE, size=10, bold=True, color="FFFFFF")
F_EXEMPLE = Font(name=POLICE, size=10, italic=True, color="7F7F7F")
FILL_ENTETE = PatternFill("solid", fgColor=BLEU_FONCE)
FILL_SAISIE = PatternFill("solid", fgColor=JAUNE_SAISIE)
FILL_CALCUL = PatternFill("solid", fgColor=GRIS_CALCUL)
FILL_REF = PatternFill("solid", fgColor=BLEU_PALE)
FIN = Side(style="thin", color="BFBFBF")
BORD = Border(left=FIN, right=FIN, top=FIN, bottom=FIN)

SAISIE_MANUELLE = {"saisie_hebdo", "sbc_activites", "sbc_synthese_as", "reunions", "points_action",
                   "sitrep_notes", "qc_suivi"}
AUTOMATIQUE = {"sitrep_meta": "s01", "epi_province": "s01", "epi_zs": "s01", "alertes_province": "s01",
               "contacts_province": "s01", "labo_province": "s01", "pec_province": "s01", "poe_province": "s01",
               "sitrep_texte": "s01", "historique": "s02, s02b", "ref_semaines": "s02", "autres_epidemies": "s02b"}

EXEMPLES = {
    "saisie_hebdo": [2026, "S41", 41, date(2026, 10, 5), date(2026, 10, 11), "ADRA", "Ituri", "Rwampara", "SURV-02",
                     "Contacts suivis chaque jour pendant les 21 jours", "", 1830, 2010, None, "ODK",
                     "Nom du point focal", date(2026, 10, 11), "Ligne d'exemple : à supprimer"],
    "sbc_activites": [date(2026, 10, 10), "ADRA", "Ituri", "Rwampara", "Dele", "Village X", "Marché", 2, 2, None,
                      None, 0, 2, 0, 0, 0, 0, 0, 0, "Signes de la MVE ; numéro vert 151", 48, 61, 1, 0, 1, 0, 1, 3,
                      "Bunia (ZS Bunia)", 1, 1, 0, "Ligne d'exemple : à supprimer"],
    "sbc_synthese_as": [date(2026, 10, 10), "ADRA", "Ituri", "Rwampara", "Dele", 25, 22, 3, 4, 40, 52, 38, 49, 1, 1, 6,
                        0, 0, 0, 1, 1, 1, 1, 2, 1, 2, 3, 2, 0, "Rumeur : « le CTE ne nourrit pas les malades »",
                        "O", "O"],
    "reunions": ["20261012-ZS-RWAMPARA", date(2026, 10, 12), "Zone de santé", "BCZ Rwampara",
                 "Réunion de coordination de la ZS", 60, 18, "Point focal ADRA", "CR_ZS_Rwampara_2026-10-12.docx", "O"],
    "points_action": ["20261012-ZS-RWAMPARA-01", "20261012-ZS-RWAMPARA", "Relancer les RECO de l'AS Dele pour la "
                      "saisie quotidienne", "O", "Superviseur ADRA", date(2026, 10, 14), "En cours", date(2026, 10, 12),
                      "Ligne d'exemple : à supprimer"],
    "sitrep_notes": [date(2026, 10, 10), "Faits marquants", 1, "Exemple : 35 RECO de l'AS de Nzere briefés sur "
                     "l'approche village.", "Nom de l'auteur"],
    "qc_suivi": ["SR-11 | SitRep 146 | Taux d'occupation", date(2026, 10, 12), "SR-11", "SitRep 146",
                 "Taux d'occupation Nord-Kivu", "84,2 % publié pour 169/209 (80,9 %)", "Majeure", "Data manager INSP",
                 "Ouverte", None],
}


def ecrire_entetes(ws, feuille):
    for j, (nom, type_, desc, oblig, valeurs) in enumerate(SCHEMA[feuille]["colonnes"], start=1):
        c = ws.cell(row=1, column=j, value=nom)
        c.font, c.fill, c.border = F_ENTETE, FILL_ENTETE, BORD
        c.alignment = Alignment(vertical="center", wrap_text=True)
        t = "formule" if type_.startswith("formule") else type_
        c.comment = Comment(f"{desc}\nType : {t}" + ("\nObligatoire" if oblig else ""), "Canevas CGPP",
                            width=260, height=90)
        largeur = max(12, min(42, len(nom) + 4))
        if t == "texte" and nom in ("texte", "action", "description", "faits_saillants", "commentaire",
                                    "libelle_fr", "libelle_en", "numerateur", "denominateur"):
            largeur = 48
        ws.column_dimensions[get_column_letter(j)].width = largeur
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 30


def listes_deroulantes(ws, feuille, refs_listes, nb_lignes=5000):
    for j, (nom, type_, _, _, valeurs) in enumerate(SCHEMA[feuille]["colonnes"], start=1):
        col = get_column_letter(j)
        plage = f"{col}2:{col}{nb_lignes}"
        dv = None
        if valeurs.startswith("liste:"):
            dv = DataValidation(type="list", formula1=refs_listes[valeurs[6:]], allow_blank=True,
                                showErrorMessage=True, errorTitle="Valeur non autorisée",
                                error="Choisir une valeur dans la liste (feuille « listes »).")
        elif nom == "zone_sante" and feuille != "ref_zones":
            dv = DataValidation(type="list", formula1="=ref_zones!$B$2:$B$600", allow_blank=True,
                                showErrorMessage=False)
        elif nom == "id_indicateur" and feuille != "ref_indicateurs":
            dv = DataValidation(type="list", formula1="=ref_indicateurs!$A$2:$A$400", allow_blank=False,
                                showErrorMessage=True, error="Identifiant absent de la matrice maîtresse.")
        elif type_ == "date":
            dv = DataValidation(type="date", operator="between", formula1="DATE(2026,1,1)",
                                formula2="DATE(2030,12,31)", allow_blank=True, showErrorMessage=True,
                                error="Date au format JJ/MM/AAAA.")
        elif type_ == "entier":
            dv = DataValidation(type="whole", operator="greaterThanOrEqual", formula1="0", allow_blank=True,
                                showErrorMessage=True, error="Nombre entier positif ou nul.")
        elif type_ == "decimal":
            dv = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True,
                                showErrorMessage=True, error="Nombre positif ou nul.")
        if dv:
            ws.add_data_validation(dv)
            dv.add(plage)


def ligne_exemple(ws, feuille):
    valeurs = EXEMPLES.get(feuille)
    if not valeurs:
        return
    cols = SCHEMA[feuille]["colonnes"]
    for j, ((nom, type_, *_), v) in enumerate(zip(cols, valeurs), start=1):
        c = ws.cell(row=2, column=j)
        if type_.startswith("formule:"):
            c.value = type_[len("formule:"):].replace("{r}", "2")
            c.fill = FILL_CALCUL
            c.number_format = "0.0"
        else:
            c.value = v
            c.fill = FILL_SAISIE
            if type_ == "date":
                c.number_format = "DD/MM/YYYY"
        c.font, c.border = F_EXEMPLE, BORD
    ws.cell(row=2, column=1).comment = Comment("Ligne d'exemple montrant le format attendu : la supprimer avant "
                                               "la première saisie réelle.", "Canevas CGPP")


def feuille_listes(wb):
    ws = wb.create_sheet("listes")
    refs = {}
    for j, (nom, valeurs) in enumerate(LISTES.items(), start=1):
        col = get_column_letter(j)
        c = ws.cell(row=1, column=j, value=nom)
        c.font, c.fill = F_ENTETE, FILL_ENTETE
        for i, v in enumerate(valeurs, start=2):
            ws.cell(row=i, column=j, value=v).font = F_TXT
        refs[nom] = f"=listes!${col}$2:${col}${len(valeurs) + 1}"
        ws.column_dimensions[col].width = max(14, max(len(str(v)) for v in valeurs) + 2)
    ws.freeze_panes = "A2"
    return refs


def feuille_dictionnaire(wb):
    ws = wb.create_sheet("dictionnaire")
    entetes = ["feuille", "titre de la feuille", "colonne", "type", "description", "obligatoire", "valeurs autorisées"]
    for j, e in enumerate(entetes, start=1):
        c = ws.cell(row=1, column=j, value=e)
        c.font, c.fill, c.border = F_ENTETE, FILL_ENTETE, BORD
    i = 2
    for feuille, spec in SCHEMA.items():
        for nom, type_, desc, oblig, valeurs in spec["colonnes"]:
            t = "formule" if type_.startswith("formule") else type_
            for j, v in enumerate([feuille, spec["titre"], nom, t, desc, "oui" if oblig else "non",
                                   valeurs.replace("liste:", "liste « ") + ("»" if valeurs.startswith("liste:") else "")],
                                  start=1):
                c = ws.cell(row=i, column=j, value=v)
                c.font, c.border = F_TXT, BORD
                c.alignment = Alignment(wrap_text=True, vertical="top")
            i += 1
    for col, w in zip("ABCDEFG", [18, 44, 30, 10, 60, 11, 24]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:G{i - 1}"


def feuille_lisezmoi(wb, cfg):
    ws = wb.active
    ws.title = "LISEZMOI"
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 30
    ws.column_dimensions["C"].width = 95
    lignes = [
        ("titre", "Base de données CGPP RDC : riposte à la MVE-17 (virus Bundibugyo)"),
        ("txt", "Canevas de la base Excel normalisée : une feuille par table, une ligne par observation, "
                "des identifiants stables. Les scripts Python du dossier « scripts » lisent et écrivent ces feuilles."),
        ("vide", ""),
        ("sous", "Légende"),
        ("legende_entete", "En-tête bleu foncé : nom technique de la colonne (ne pas modifier). Survoler l'en-tête "
                           "pour lire la description."),
        ("legende_saisie", "Fond jaune : cellule à saisir (ligne d'exemple en italique, à supprimer avant usage)."),
        ("legende_calcul", "Fond gris : cellule calculée par formule (ne pas saisir)."),
        ("legende_ref", "Fond bleu pâle : référentiel tenu par l'équipe SI (zones, indicateurs, semaines)."),
        ("vide", ""),
        ("sous", "Qui remplit quoi"),
        ("kv", ("saisie_hebdo", "Points focaux des IP, chaque vendredi (dimanche au plus tard) : une ligne par "
                                "indicateur, ZS et semaine ; numérateur ET dénominateur pour les pourcentages.")),
        ("kv", ("sbc_activites / sbc_synthese_as", "Export de la fiche SBC journalière (ODK), une ligne par site "
                                                   "visité et une ligne par aire de santé et par jour.")),
        ("kv", ("reunions / points_action", "Représentant CGPP à chaque réunion : la réunion, puis chaque point "
                                            "d'action avec responsable, échéance et statut.")),
        ("kv", ("sitrep_notes", "Équipe SI : faits marquants, défis et actions repris dans le SitRep CGPP.")),
        ("kv", ("feuilles sitrep_*, epi_*, …", "Remplies automatiquement par le script s01 à partir des SitReps "
                                                "nationaux : ne pas saisir à la main.")),
        ("kv", ("tableau_suivi", "Calculé par formules à partir de saisie_hebdo et ref_indicateurs.")),
        ("kv", ("qc_suivi", "Alimentée par le contrôle qualité (s03) ; l'équipe complète responsable, statut et "
                            "date de résolution.")),
        ("vide", ""),
        ("sous", "Règles de saisie"),
        ("txt", "1. Une valeur par cellule ; pas de texte dans les colonnes numériques ; pas de cellules fusionnées."),
        ("txt", "2. Les codes (province, ZS, partenaire, indicateur) viennent des listes déroulantes."),
        ("txt", "3. Un pourcentage se saisit toujours en numérateur et dénominateur, jamais en pourcentage tout fait."),
        ("txt", "4. Une semaine se renseigne même si la valeur est zéro : une cellule vide signifie « non rapporté »."),
        ("vide", ""),
        ("sous", "Commandes"),
        ("txt", "python pipeline.py            (extraction, contrôle qualité, analyses et SitRep CGPP)"),
        ("txt", "python pipeline.py --etapes qc,sitrep   (relancer seulement le contrôle et le SitRep)"),
        ("vide", ""),
        ("txt", f"Canevas généré le {datetime.now():%d/%m/%Y} ; modèle de données : {len(SCHEMA)} tables."),
    ]
    r = 1
    for type_, contenu in lignes:
        if type_ == "titre":
            ws.cell(row=r, column=2, value=contenu).font = F_TITRE
        elif type_ == "sous":
            ws.cell(row=r, column=2, value=contenu).font = F_SOUS
        elif type_ == "txt":
            c = ws.cell(row=r, column=2, value=contenu)
            c.font = F_TXT
        elif type_.startswith("legende"):
            fill = {"legende_entete": FILL_ENTETE, "legende_saisie": FILL_SAISIE, "legende_calcul": FILL_CALCUL,
                    "legende_ref": FILL_REF}[type_]
            ws.cell(row=r, column=2).fill = fill
            ws.cell(row=r, column=2).border = BORD
            ws.cell(row=r, column=3, value=contenu).font = F_TXT
        elif type_ == "kv":
            ws.cell(row=r, column=2, value=contenu[0]).font = Font(name=POLICE, size=10, bold=True)
            c = ws.cell(row=r, column=3, value=contenu[1])
            c.font = F_TXT
            c.alignment = Alignment(wrap_text=True, vertical="top")
        r += 1
    ws.sheet_view.showGridLines = False


def feuille_tableau_suivi(wb, nb=220):
    ws = wb.create_sheet("tableau_suivi")
    ri = {c: get_column_letter(i + 1) for i, c in enumerate(colonnes("ref_indicateurs"))}
    sh = {c: get_column_letter(i + 1) for i, c in enumerate(colonnes("saisie_hebdo"))}
    ws["A1"] = "Suivi des indicateurs : calculé à partir de la feuille saisie_hebdo"
    ws["A1"].font = F_TITRE
    ws["A2"] = ("Les pourcentages sont recalculés à partir de la somme des numérateurs et de la somme des "
                "dénominateurs (jamais en additionnant des pourcentages). Ne rien saisir dans cette feuille.")
    ws["A2"].font = F_TXT
    entetes = ["id_indicateur", "pilier", "indicateur", "type", "cible DOS", "cible DOS (nombre)",
               "numérateur cumulé", "dénominateur cumulé", "réalisé", "% de la cible", "dernière semaine saisie",
               "lignes saisies", "statut"]
    for j, e in enumerate(entetes, start=1):
        c = ws.cell(row=4, column=j, value=e)
        c.font, c.fill, c.border = F_ENTETE, FILL_ENTETE, BORD
        c.alignment = Alignment(wrap_text=True, vertical="center")
    R_ = "ref_indicateurs!"
    S_ = "saisie_hebdo!"
    for i in range(5, 5 + nb):
        k = i - 3   # rang dans ref_indicateurs (ligne 2 = premier indicateur)
        id_ = f"$A{i}"
        cherche = lambda col: f'IF({id_}="","",INDEX({R_}${ri[col]}:${ri[col]},MATCH({id_},{R_}$A:$A,0))&"")'
        f = {
            1: f'=IFERROR(INDEX({R_}$A:$A,{k})&"","")',
            2: "=" + cherche("pilier"),
            3: "=" + cherche("libelle_fr"),
            4: "=" + cherche("type"),
            5: "=" + cherche("cible_dos"),
            6: f'=IF({id_}="","",IF(INDEX({R_}${ri["cible_dos_num"]}:${ri["cible_dos_num"]},'
               f'MATCH({id_},{R_}$A:$A,0))="","",INDEX({R_}${ri["cible_dos_num"]}:${ri["cible_dos_num"]},'
               f'MATCH({id_},{R_}$A:$A,0))))',
            7: f'=IF({id_}="","",SUMIFS({S_}${sh["numerateur"]}:${sh["numerateur"]},{S_}${sh["id_indicateur"]}:'
               f'${sh["id_indicateur"]},{id_}))',
            8: f'=IF({id_}="","",SUMIFS({S_}${sh["denominateur"]}:${sh["denominateur"]},{S_}${sh["id_indicateur"]}:'
               f'${sh["id_indicateur"]},{id_}))',
            12: f'=IF({id_}="","",COUNTIFS({S_}${sh["id_indicateur"]}:${sh["id_indicateur"]},{id_}))',
            9: f'=IF(OR({id_}="",L{i}=0),"",IF(D{i}="Pourcentage",IF(H{i}>0,G{i}/H{i}*100,""),G{i}))',
            10: f'=IF(OR(I{i}="",F{i}="",F{i}=0),"",I{i}/F{i})',
            11: f'=IF(OR({id_}="",L{i}=0),"",IF(_xlfn.MAXIFS({S_}${sh["semaine_iso"]}:${sh["semaine_iso"]},'
                f'{S_}${sh["id_indicateur"]}:${sh["id_indicateur"]},{id_})=0,"",_xlfn.MAXIFS({S_}${sh["semaine_iso"]}:'
                f'${sh["semaine_iso"]},{S_}${sh["id_indicateur"]}:${sh["id_indicateur"]},{id_})))',
            13: f'=IF({id_}="","",IF(L{i}=0,"Non rapporté",IF(F{i}="","Suivi sans cible",IF(D{i}="Pourcentage",'
                f'IF(I{i}>=F{i},"Cible atteinte","Sous la cible"),IF(J{i}>=1,"Cible atteinte","En cours")))))',
        }
        for j, formule in f.items():
            c = ws.cell(row=i, column=j, value=formule)
            c.font, c.fill, c.border = F_TXT, FILL_CALCUL, BORD
        ws.cell(row=i, column=9).number_format = "#,##0.0"
        ws.cell(row=i, column=10).number_format = "0%"
        ws.cell(row=i, column=7).number_format = "#,##0"
        ws.cell(row=i, column=8).number_format = "#,##0"
    for col, w in zip("ABCDEFGHIJKLM", [14, 9, 58, 12, 16, 12, 13, 13, 11, 11, 12, 10, 16]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "D5"
    ws.row_dimensions[4].height = 32


def construire(cfg: dict) -> Path:
    wb = openpyxl.Workbook()
    feuille_lisezmoi(wb, cfg)
    refs = feuille_listes(wb)
    for feuille, spec in SCHEMA.items():
        ws = wb.create_sheet(feuille)
        ecrire_entetes(ws, feuille)
        listes_deroulantes(ws, feuille, refs)
        if feuille in SAISIE_MANUELLE:
            ligne_exemple(ws, feuille)
            ws.sheet_properties.tabColor = "F2C200"
        elif feuille.startswith("ref_"):
            ws.sheet_properties.tabColor = "8EA9DB"
        else:
            ws.sheet_properties.tabColor = "A6A6A6"
    # référentiels pré-remplis
    ws = wb["ref_indicateurs"]
    m = matrice().assign(cible_dos_num=lambda d: d.cible_dos.map(cible_numerique),
                         cible_t5_num=lambda d: d.cible_tranche5.map(cible_numerique))
    for i, ligne in enumerate(m.reindex(columns=colonnes("ref_indicateurs")).itertuples(index=False), start=2):
        for j, v in enumerate(ligne, start=1):
            if v is not None and v == v and v != "":
                c = ws.cell(row=i, column=j, value=v)
                c.font, c.fill = F_TXT, FILL_REF
    ws = wb["ref_partenaires"]
    for i, (k, v) in enumerate(cfg["partenaires"].items(), start=2):
        for j, val in enumerate([k, v, "Partenaire de mise en oeuvre"], start=1):
            c = ws.cell(row=i, column=j, value=val)
            c.font, c.fill = F_TXT, FILL_REF
    ws = wb["ref_zones"]
    for i, (zs, info) in enumerate(cfg["zones_cgpp"].items(), start=2):
        for j, val in enumerate([info["province"], zs, None, info["ip"], "intervention", None, "config.yaml"], start=1):
            if val is not None:
                c = ws.cell(row=i, column=j, value=val)
                c.font, c.fill = F_TXT, FILL_REF
    feuille_tableau_suivi(wb)
    feuille_dictionnaire(wb)
    ordre = (["LISEZMOI", "tableau_suivi"] + [f for f in SCHEMA if f in SAISIE_MANUELLE] +
             [f for f in SCHEMA if f.startswith("ref_")] +
             [f for f in SCHEMA if f not in SAISIE_MANUELLE and not f.startswith("ref_")] + ["dictionnaire", "listes"])
    wb._sheets = [wb[n] for n in ordre]
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if c.font is None or c.font.name != POLICE:
                    c.font = Font(name=POLICE, size=c.font.size if c.font else 10, bold=c.font.bold if c.font else False,
                                  italic=c.font.italic if c.font else False,
                                  color=c.font.color if c.font else None)
    wb.calculation.fullCalcOnLoad = True
    sortie = cfg["chemins"]["canevas_base"]
    sortie.parent.mkdir(parents=True, exist_ok=True)
    wb.save(sortie)
    LOG.info("Canevas écrit : %s (%d feuilles)", sortie, len(wb.sheetnames))
    return sortie


def main(cfg: dict | None = None):
    cfg = cfg or charger_config()
    return construire(cfg)


if __name__ == "__main__":
    main()
