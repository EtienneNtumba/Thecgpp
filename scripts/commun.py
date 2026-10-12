"""
commun.py - fonctions partagées de la chaîne SI CGPP (riposte MVE-17).

Contient :
  - le chargement de la configuration (config.yaml) ;
  - le MODELE DE DONNEES de la base Excel (SCHEMA), source unique pour le
    canevas, l'écriture de la base et le dictionnaire de données ;
  - les utilitaires de lecture des nombres au format français, de
    normalisation des provinces et des zones de santé ;
  - la conversion .doc -> .docx via LibreOffice ;
  - l'écriture d'une feuille de la base en conservant la mise en forme.
"""
from __future__ import annotations

import logging
import re
import shutil
import subprocess
import tempfile
import unicodedata
import warnings
from pathlib import Path

import pandas as pd
import yaml

RACINE = Path(__file__).resolve().parents[1]

# --------------------------------------------------------------------------
# Configuration et journalisation
# --------------------------------------------------------------------------

def charger_config(chemin: str | Path | None = None) -> dict:
    chemin = Path(chemin) if chemin else RACINE / "config.yaml"
    with open(chemin, encoding="utf-8") as f:
        cfg = yaml.safe_load(f)
    cfg["chemins"] = {k: (RACINE / v) for k, v in cfg["chemins"].items()}
    for cle in ("sorties_qc", "sorties_analyses", "sorties_figures", "sorties_sitrep",
                "intermediaire", "referentiels"):
        cfg["chemins"][cle].mkdir(parents=True, exist_ok=True)
    cfg["chemins"]["base"].parent.mkdir(parents=True, exist_ok=True)
    return cfg


def journal(nom: str) -> logging.Logger:
    """Journal nommé : messages à l'écran (le pipeline ajoute en plus un fichier journal)."""
    log = logging.getLogger(nom)
    log.setLevel(logging.INFO)
    log.propagate = False
    if not any(type(h) is logging.StreamHandler for h in log.handlers):  # un FileHandler ne compte pas
        h = logging.StreamHandler()
        h.setFormatter(logging.Formatter("%(asctime)s | %(name)s | %(levelname)s | %(message)s",
                                         "%H:%M:%S"))
        log.addHandler(h)
    return log


# --------------------------------------------------------------------------
# Référentiels géographiques
# --------------------------------------------------------------------------

PROVINCES = ["Ituri", "Nord-Kivu", "Haut-Uélé", "Bas-Uélé", "Tshopo", "Sud-Kivu", "Sud-Ubangi"]
CODES_PROVINCE = {"IT": "Ituri", "NK": "Nord-Kivu", "SK": "Sud-Kivu", "HU": "Haut-Uélé",
                  "BU": "Bas-Uélé", "TS": "Tshopo", "SU": "Sud-Ubangi"}


def sans_accents(s: str) -> str:
    return "".join(c for c in unicodedata.normalize("NFD", s) if unicodedata.category(c) != "Mn")


def cle(s) -> str:
    """Clé de comparaison : minuscules, sans accents, sans espaces ni tirets."""
    if s is None or (isinstance(s, float) and pd.isna(s)):
        return ""
    return re.sub(r"[\s\-_'’.]+", "", sans_accents(str(s)).lower())


_PROV_PAR_CLE = {cle(p): p for p in PROVINCES}
_PROV_PAR_CLE.update({"nordkivi": "Nord-Kivu", "hautuele": "Haut-Uélé", "basuele": "Bas-Uélé",
                      "northkivu": "Nord-Kivu", "southkivu": "Sud-Kivu", "southubangi": "Sud-Ubangi",
                      "upperuele": "Haut-Uélé", "loweruele": "Bas-Uélé"})


def norm_province(s) -> str | None:
    if s is None:
        return None
    k = cle(s)
    if k in _PROV_PAR_CLE:
        return _PROV_PAR_CLE[k]
    if str(s).strip().upper() in CODES_PROVINCE:
        return CODES_PROVINCE[str(s).strip().upper()]
    return None


def norm_zs(s, referentiel: dict | None = None) -> str:
    """Nom canonique d'une zone de santé. `referentiel` : {cle(nom): nom canonique}."""
    brut = re.sub(r"\s+", " ", str(s)).strip()
    if referentiel and cle(brut) in referentiel:
        return referentiel[cle(brut)]
    # Mise en forme par défaut : première lettre de chaque élément en capitale
    return "-".join(p.capitalize() if p.isupper() or p.islower() else p for p in brut.split("-"))


# --------------------------------------------------------------------------
# Nombres au format français
# --------------------------------------------------------------------------

_ESPACES = "    "


def nombre_fr(s):
    """'3 025' -> 3025.0 ; '46,2%' -> 46.2 ; 'NA'/'ND'/'' -> None."""
    if s is None:
        return None
    if isinstance(s, (int, float)):
        return None if pd.isna(s) else float(s)
    t = str(s).strip().replace("**", "")
    if t.upper() in {"", "NA", "ND", "N/A", "-", "—"}:
        return None
    t = t.replace("%", "")
    for e in _ESPACES:
        t = t.replace(e, "")
    t = t.replace(",", ".")
    m = re.match(r"^-?\d+(\.\d+)?$", t)
    return float(t) if m else None


def fraction_fr(s):
    """'28/36 (77,8 %)' -> (28, 36, 77.8) ; '93,3% (56/60)' -> (56, 60, 93.3)."""
    if s is None:
        return (None, None, None)
    t = str(s)
    for e in _ESPACES:
        t = t.replace(e, "")
    num = den = pct = None
    m = re.search(r"(\d+)/(\d+)", t)
    if m:
        num, den = float(m.group(1)), float(m.group(2))
    m2 = re.search(r"(\d+(?:[.,]\d+)?)%", t)
    if m2:
        pct = float(m2.group(1).replace(",", "."))
    return (num, den, pct)


def pourcentage(num, den, decimales: int = 1):
    if num is None or den in (None, 0) or pd.isna(num) or pd.isna(den):
        return None
    return round(100.0 * float(num) / float(den), decimales)


# --------------------------------------------------------------------------
# Conversion de documents Word
# --------------------------------------------------------------------------

EMPLACEMENTS_SOFFICE = [
    "C:/Program Files/LibreOffice/program/soffice.exe",
    "C:/Program Files (x86)/LibreOffice/program/soffice.exe",
    "/Applications/LibreOffice.app/Contents/MacOS/soffice",
    "/usr/bin/soffice", "/usr/local/bin/soffice", "/opt/libreoffice/program/soffice",
]


def trouver_soffice(soffice: str = "soffice") -> str:
    """Chemin de LibreOffice : valeur de config.yaml, sinon PATH, sinon emplacements usuels."""
    for candidat in [soffice, "soffice", "libreoffice"]:
        if candidat and (Path(candidat).is_file() or shutil.which(candidat)):
            return shutil.which(candidat) or candidat
    for candidat in EMPLACEMENTS_SOFFICE:
        if Path(candidat).is_file():
            return candidat
    raise FileNotFoundError("LibreOffice introuvable : l'installer (https://www.libreoffice.org) ou indiquer "
                            "le chemin de soffice dans config.yaml (clé libreoffice).")


def convertir(chemin: Path, format_cible: str, dossier_sortie: Path, soffice: str = "soffice") -> Path:
    """Conversion LibreOffice sans interface, avec un profil dédié : fonctionne même si LibreOffice
    est déjà ouvert sur le poste."""
    dossier_sortie.mkdir(parents=True, exist_ok=True)
    profil = Path(tempfile.gettempdir()) / "lo_profil_cgpp_si"
    cmd = [trouver_soffice(soffice), f"-env:UserInstallation={profil.as_uri()}", "--headless",
           "--convert-to", format_cible, "--outdir", str(dossier_sortie), str(chemin)]
    subprocess.run(cmd, check=True, capture_output=True, timeout=300)
    cible = dossier_sortie / (Path(chemin).stem + "." + format_cible)
    if not cible.exists():
        raise RuntimeError(f"Conversion LibreOffice échouée : {chemin} vers {format_cible}")
    return cible


def vers_docx(chemin: Path, dossier_sortie: Path, soffice: str = "soffice") -> Path:
    """Convertit un .doc en .docx (LibreOffice). Renvoie le chemin .docx (cache)."""
    chemin = Path(chemin)
    if chemin.suffix.lower() == ".docx":
        return chemin
    cible = dossier_sortie / (chemin.stem + ".docx")
    if cible.exists() and cible.stat().st_mtime >= chemin.stat().st_mtime:
        return cible
    return convertir(chemin, "docx", dossier_sortie, soffice)


def vers_pdf(chemin_docx: Path, dossier_sortie: Path, soffice: str = "soffice") -> Path:
    return convertir(chemin_docx, "pdf", dossier_sortie, soffice)


# --------------------------------------------------------------------------
# MODELE DE DONNEES (une feuille Excel = une table)
# Chaque colonne : (nom, type, description, obligatoire, valeurs autorisées)
# --------------------------------------------------------------------------

TYPES_SITE = ["CPN", "CPS", "Porte à porte (VAD)", "Lieu religieux (église)", "Chambre de prière",
              "École", "Marché", "Rond-point", "Radio", "Site de guérison traditionnelle",
              "Ferme laitière / centre d'engraissement", "Point d'eau",
              "Site de distribution alimentaire ou cash", "Camp de réfugiés", "Site de déplacés",
              "Point d'entrée", "Boucherie", "Restaurant", "Carré / site minier",
              "Parking véhicules", "Parking motos", "Autre"]

SCHEMA: dict[str, dict] = {
    "ref_zones": {
        "titre": "Référentiel des zones de santé",
        "cle": ["province", "zone_sante"],
        "colonnes": [
            ("province", "texte", "Province (nom canonique)", True, "liste:provinces"),
            ("zone_sante", "texte", "Zone de santé (nom canonique du SitRep national)", True, ""),
            ("code_zs", "texte", "Code ou UID DHIS2 de la ZS (à compléter)", False, ""),
            ("ip_cgpp", "texte", "Partenaire CGPP intervenant dans la ZS", False, "liste:ip"),
            ("statut_cgpp", "texte", "Intervention / préparation / hors zone CGPP", False,
             "liste:statut_cgpp"),
            ("population", "entier", "Population de la ZS (dénominateur des taux)", False, ""),
            ("source", "texte", "Origine de la ligne", False, ""),
        ],
    },
    "ref_partenaires": {
        "titre": "Référentiel des partenaires de mise en oeuvre",
        "cle": ["code_ip"],
        "colonnes": [
            ("code_ip", "texte", "Code court du partenaire", True, ""),
            ("nom", "texte", "Nom complet", True, ""),
            ("role", "texte", "Rôle dans le consortium", False, ""),
            ("provinces", "texte", "Provinces couvertes", False, ""),
            ("point_focal", "texte", "Point focal données", False, ""),
            ("email_point_focal", "texte", "Courriel du point focal", False, ""),
        ],
    },
    "ref_indicateurs": {
        "titre": "Matrice maîtresse des indicateurs",
        "cle": ["id_indicateur"],
        "colonnes": [
            ("id_indicateur", "texte", "Identifiant stable (texte, ex. SURV.05)", True, ""),
            ("pilier", "texte", "Pilier du plan de travail DOS", True, "liste:piliers"),
            ("libelle_fr", "texte", "Intitulé en français", True, ""),
            ("libelle_en", "texte", "Intitulé en anglais (rapport DOS)", False, ""),
            ("type", "texte", "Nombre / Pourcentage / Délai", True, "liste:type_indicateur"),
            ("niveau", "texte", "Tier 1 (performance) / Tier 2 (extrant) / Contexte", False, ""),
            ("numerateur", "texte", "Définition du numérateur", True, ""),
            ("denominateur", "texte", "Définition du dénominateur (N/A pour un nombre)", False, ""),
            ("source_donnees", "texte", "Source primaire", True, ""),
            ("champ_source", "texte", "Champ(s) de la fiche SBC / ODK / SitRep", False, ""),
            ("frequence", "texte", "Fréquence de rapportage", True, ""),
            ("desagregation", "texte", "Désagrégations requises", False, ""),
            ("cible_dos", "texte", "Cible du plan de travail DOS", False, ""),
            ("cible_tranche5", "texte", "Cible de la Tranche 5", False, ""),
            ("cible_dos_num", "decimal", "Cible DOS en nombre (pour les calculs ; % en points)", False, ""),
            ("cible_t5_num", "decimal", "Cible Tranche 5 en nombre", False, ""),
            ("seuil_alerte", "texte", "Seuil de performance (code couleur)", False, ""),
            ("id_barclay", "texte", "Numéro dans la matrice de Dr Barclay", False, ""),
            ("id_cadre_global", "texte", "Rang dans le cadre global (1 à 25)", False, ""),
            ("statut", "texte", "existant / à corriger / nouveau", True, "liste:statut_indicateur"),
            ("priorite", "texte", "Priorité de mise en oeuvre (1 à 3)", False, ""),
            ("commentaire", "texte", "Remarques, écarts relevés", False, ""),
        ],
    },
    "ref_semaines": {
        "titre": "Calendrier des semaines épidémiologiques (ISO)",
        "cle": ["annee", "semaine_iso"],
        "colonnes": [
            ("annee", "entier", "Année ISO", True, ""),
            ("semaine_iso", "entier", "Numéro de semaine ISO", True, ""),
            ("date_debut", "date", "Lundi", True, ""),
            ("date_fin", "date", "Dimanche", True, ""),
            ("libelle", "texte", "Libellé affiché", True, ""),
        ],
    },
    "sitrep_meta": {
        "titre": "SitReps nationaux : en-tête et chiffres clés",
        "cle": ["sitrep_no"],
        "colonnes": [
            ("sitrep_no", "entier", "Numéro du SitRep national", True, ""),
            ("date_rapport", "date", "Date des données", True, ""),
            ("date_publication", "date", "Date de publication", False, ""),
            ("fichier_source", "texte", "Fichier d'origine", True, ""),
            ("cas_cumul", "entier", "Cumul des cas confirmés (encadré)", True, ""),
            ("deces_cumul", "entier", "Cumul des décès confirmés (encadré)", True, ""),
            ("gueris_cumul", "entier", "Cumul des guéris (encadré)", False, ""),
            ("taux_suivi_contacts", "decimal", "Taux de suivi des contacts du jour (%)", False, ""),
            ("letalite", "decimal", "Létalité publiée (%)", False, ""),
            ("patients_isolement", "entier", "Patients en isolement / CTE (encadré)", False, ""),
            ("aires_sante_touchees", "entier", "Aires de santé touchées", False, ""),
            ("aires_sante_total", "entier", "Total des aires de santé", False, ""),
            ("nouveaux_cas_24h", "entier", "Nouveaux cas confirmés (faits saillants)", False, ""),
            ("deces_communautaires_24h", "entier", "Décès communautaires du jour", False, ""),
            ("deces_cte_24h", "entier", "Décès dans les CTE du jour", False, ""),
            ("gueris_24h", "entier", "Guéris du jour (résumé)", False, ""),
            ("echantillons_24h", "entier", "Échantillons analysés (faits saillants)", False, ""),
            ("faits_saillants", "texte", "Texte des faits saillants", False, ""),
            ("date_extraction", "date", "Date de l'extraction", True, ""),
        ],
    },
    "epi_province": {
        "titre": "Situation par province (Tableau 1 du SitRep)",
        "cle": ["sitrep_no", "province"],
        "colonnes": [
            ("sitrep_no", "entier", "Numéro du SitRep", True, ""),
            ("date_rapport", "date", "Date des données", True, ""),
            ("province", "texte", "Province", True, "liste:provinces"),
            ("nouveaux_cas_24h", "entier", "Nouveaux cas confirmés (24 h)", True, ""),
            ("cas_cumul", "entier", "Cas confirmés cumulés", True, ""),
            ("deces_cumul", "entier", "Décès confirmés cumulés", True, ""),
            ("letalite_publiee", "decimal", "Létalité publiée (%)", False, ""),
            ("zs_touchees", "entier", "ZS touchées", False, ""),
            ("zs_total", "entier", "ZS de la province", False, ""),
        ],
    },
    "epi_zs": {
        "titre": "Situation par zone de santé (Tableau 2 du SitRep)",
        "cle": ["sitrep_no", "province", "zone_sante"],
        "colonnes": [
            ("sitrep_no", "entier", "Numéro du SitRep", True, ""),
            ("date_rapport", "date", "Date des données", True, ""),
            ("province", "texte", "Province", True, "liste:provinces"),
            ("zone_sante", "texte", "Zone de santé", True, ""),
            ("cas_cumul", "entier", "Cas confirmés cumulés", False, ""),
            ("deces_cumul", "entier", "Décès confirmés cumulés", True, ""),
            ("letalite_publiee", "decimal", "Létalité publiée (%)", False, ""),
            ("nouveaux_cas_24h", "entier", "Nouveaux cas (24 h)", False, ""),
            ("deces_communautaires_24h", "entier", "Décès communautaires swab+ (24 h)", False, ""),
            ("deces_cte_24h", "entier", "Décès confirmés intra-CTE (24 h)", False, ""),
            ("deces_total_24h", "entier", "Total des décès confirmés (24 h)", False, ""),
            ("est_a_ventiler", "booleen", "Ligne « à ventiler » (décès non attribués à une ZS)",
             False, ""),
        ],
    },
    "alertes_province": {
        "titre": "Alertes par province (Tableau 3 du SitRep)",
        "cle": ["sitrep_no", "province"],
        "colonnes": [
            ("sitrep_no", "entier", "Numéro du SitRep", True, ""),
            ("date_rapport", "date", "Date des données", True, ""),
            ("province", "texte", "Province", True, "liste:provinces"),
            ("statut_donnees", "texte", "OK ou ND (non disponible)", True, "liste:statut_donnees"),
            ("recues_vivants", "entier", "Alertes reçues : vivants", False, ""),
            ("recues_deces", "entier", "Alertes reçues : décès", False, ""),
            ("recues_total", "entier", "Total des alertes", False, ""),
            ("validees_vivants", "entier", "Validées (cas suspects) : vivants", False, ""),
            ("validees_deces", "entier", "Validées : décès", False, ""),
            ("invalidees_vivants", "entier", "Invalidées : vivants", False, ""),
            ("invalidees_deces", "entier", "Invalidées : décès", False, ""),
            ("suspects_investigues", "entier", "Suspects investigués", False, ""),
            ("suspects_transferes", "entier", "Suspects transférés au CT/CTE", False, ""),
        ],
    },
    "contacts_province": {
        "titre": "Suivi des contacts par province",
        "cle": ["sitrep_no", "province"],
        "colonnes": [
            ("sitrep_no", "entier", "Numéro du SitRep", True, ""),
            ("date_rapport", "date", "Date des données", True, ""),
            ("province", "texte", "Province", True, "liste:provinces"),
            ("statut_donnees", "texte", "OK ou ND", True, "liste:statut_donnees"),
            ("contacts_suivis", "entier", "Contacts en cours de suivi", False, ""),
            ("contacts_vus", "entier", "Contacts vus dans les 24 h", False, ""),
            ("taux_publie", "decimal", "Taux publié (%)", False, ""),
        ],
    },
    "labo_province": {
        "titre": "Laboratoire par province",
        "cle": ["sitrep_no", "province"],
        "colonnes": [
            ("sitrep_no", "entier", "Numéro du SitRep", True, ""),
            ("date_rapport", "date", "Date des données", True, ""),
            ("province", "texte", "Province", True, "liste:provinces"),
            ("echantillons", "entier", "Échantillons reçus et analysés", False, ""),
            ("positifs", "entier", "Résultats positifs", False, ""),
            ("positifs_vivants", "entier", "Positifs vivants", False, ""),
            ("positifs_deces", "entier", "Positifs décédés", False, ""),
            ("positivite_publiee", "decimal", "Positivité publiée (%)", False, ""),
            ("texte_source", "texte", "Phrase source", False, ""),
        ],
    },
    "pec_province": {
        "titre": "Prise en charge par province",
        "cle": ["sitrep_no", "province"],
        "colonnes": [
            ("sitrep_no", "entier", "Numéro du SitRep", True, ""),
            ("date_rapport", "date", "Date des données", True, ""),
            ("province", "texte", "Province", True, "liste:provinces"),
            ("admissions", "entier", "Nouvelles admissions", False, ""),
            ("sorties", "entier", "Sorties", False, ""),
            ("guerisons", "entier", "Guérisons", False, ""),
            ("patients_isolement", "entier", "Patients en isolement en fin de journée", False, ""),
            ("texte_source", "texte", "Phrase source", False, ""),
        ],
    },
    "poe_province": {
        "titre": "Indicateurs PoE/PoC (Tableau 4 du SitRep)",
        "cle": ["sitrep_no", "indicateur", "province"],
        "colonnes": [
            ("sitrep_no", "entier", "Numéro du SitRep", True, ""),
            ("date_rapport", "date", "Date des données", True, ""),
            ("indicateur", "texte", "Libellé de l'indicateur", True, ""),
            ("province", "texte", "Province ou Global", True, ""),
            ("valeur_texte", "texte", "Valeur telle que publiée", False, ""),
            ("valeur", "decimal", "Valeur numérique", False, ""),
            ("numerateur", "decimal", "Numérateur (si publié)", False, ""),
            ("denominateur", "decimal", "Dénominateur (si publié)", False, ""),
        ],
    },
    "sitrep_texte": {
        "titre": "Texte des SitReps par section (pour recherche et contrôle)",
        "cle": ["sitrep_no", "ordre"],
        "colonnes": [
            ("sitrep_no", "entier", "Numéro du SitRep", True, ""),
            ("ordre", "entier", "Ordre du paragraphe", True, ""),
            ("section", "texte", "Pilier ou section", False, ""),
            ("sous_section", "texte", "Principales actions / Défis", False, ""),
            ("texte", "texte", "Paragraphe", True, ""),
        ],
    },
    "historique": {
        "titre": "Points de référence issus d'autres documents",
        "cle": ["date_reference", "source", "territoire", "indicateur"],
        "colonnes": [
            ("date_reference", "date", "Date à laquelle se rapporte la valeur", True, ""),
            ("source", "texte", "Document d'origine", True, ""),
            ("territoire", "texte", "National ou province", True, ""),
            ("indicateur", "texte", "Indicateur", True, ""),
            ("valeur", "decimal", "Valeur", True, ""),
            ("commentaire", "texte", "Remarque", False, ""),
        ],
    },
    "saisie_hebdo": {
        "titre": "Saisie hebdomadaire des IP (format long : une ligne = un indicateur, une ZS, une semaine)",
        "cle": ["periode", "ip", "province", "zone_sante", "id_indicateur", "desagregation"],
        "colonnes": [
            ("annee", "entier", "Année", True, ""),
            ("periode", "texte", "S36, S37... ou période agrégée (S20-S35)", True, ""),
            ("semaine_iso", "entier", "Numéro de semaine ISO (vide si période agrégée)", False, ""),
            ("date_debut", "date", "Premier jour de la période", True, ""),
            ("date_fin", "date", "Dernier jour de la période", True, ""),
            ("ip", "texte", "Partenaire", True, "liste:ip"),
            ("province", "texte", "Province", True, "liste:provinces"),
            ("zone_sante", "texte", "Zone de santé (vide si non ventilé)", False, ""),
            ("id_indicateur", "texte", "Identifiant de la matrice maîtresse", True, ""),
            ("libelle_indicateur", "texte", "Rappel de l'intitulé", False, ""),
            ("desagregation", "texte", "Catégorie de désagrégation (ex. canal=VAD, sexe=F)", False, ""),
            ("numerateur", "decimal", "Numérateur (ou valeur pour un nombre)", True, ""),
            ("denominateur", "decimal", "Dénominateur (pourcentages uniquement)", False, ""),
            ("valeur", 'formule:=IF(M{r}="",L{r},IF(M{r}=0,"",L{r}/M{r}*100))',
             "Valeur calculée par formule (numérateur/dénominateur x 100, ou numérateur)", False, ""),
            ("source_saisie", "texte", "ODK / registre / classeur historique", False, ""),
            ("saisi_par", "texte", "Nom de la personne", False, ""),
            ("date_saisie", "date", "Date de saisie", False, ""),
            ("commentaire", "texte", "Remarque", False, ""),
        ],
    },
    "sbc_activites": {
        "titre": "Fiche SBC journalière : activités par type de site (format long)",
        "cle": ["date", "aire_sante", "village_quartier", "type_site"],
        "colonnes": [
            ("date", "date", "Date de l'activité", True, ""),
            ("ip", "texte", "Partenaire", True, "liste:ip"),
            ("province", "texte", "Province", True, "liste:provinces"),
            ("zone_sante", "texte", "Zone de santé", True, ""),
            ("aire_sante", "texte", "Aire de santé", True, ""),
            ("village_quartier", "texte", "Village / quartier / CAC", False, ""),
            ("type_site", "texte", "Type de site visité", True, "liste:types_site"),
            ("nb_seances_prevues", "entier", "Séances prévues", False, ""),
            ("nb_seances_realisees", "entier", "Séances réalisées", False, ""),
            ("nb_menages_prevus", "entier", "Ménages prévus (porte à porte uniquement)", False, ""),
            ("nb_menages_visites", "entier", "Ménages visités (porte à porte uniquement)", False, ""),
            ("nb_vad", "entier", "Visites à domicile", False, ""),
            ("nb_causeries", "entier", "Causeries éducatives", False, ""),
            ("nb_ecoute_negociation", "entier", "Écoute active / négociation", False, ""),
            ("nb_debats_dialogues", "entier", "Débats / dialogues communautaires", False, ""),
            ("nb_sketchs", "entier", "Sketchs", False, ""),
            ("nb_emissions_radio", "entier", "Émissions radio", False, ""),
            ("nb_emissions_interactives", "entier", "Émissions interactives", False, ""),
            ("nb_spots_radio", "entier", "Spots radio", False, ""),
            ("themes", "texte", "Thèmes abordés", False, ""),
            ("sensibilises_h", "entier", "Personnes sensibilisées : hommes", False, ""),
            ("sensibilises_f", "entier", "Personnes sensibilisées : femmes", False, ""),
            ("alertes_malades", "entier", "Alertes MVE trouvées : malades", False, ""),
            ("alertes_deces", "entier", "Alertes MVE trouvées : décès", False, ""),
            ("alertes_malades_notifiees_24h", "entier", "Alertes malades signalées au CS en 24 h",
             False, ""),
            ("alertes_deces_notifiees_24h", "entier", "Alertes décès signalées au CS en 24 h",
             False, ""),
            ("alertes_investiguees_24h", "entier", "Alertes investiguées en 24 h", False, ""),
            ("nb_visiteurs", "entier", "Visiteurs", False, ""),
            ("provenance_visiteurs", "texte", "Village et ZS de provenance", False, ""),
            ("malades_orientes_fosa", "entier", "Malades orientés vers une FOSA", False, ""),
            ("malades_arrives_fosa", "entier", "Malades arrivés à la FOSA", False, ""),
            ("malades_orientes_ct_cte", "entier", "Malades orientés vers CT/CTE/isolement", False, ""),
            ("commentaire", "texte", "Remarque", False, ""),
        ],
    },
    "sbc_synthese_as": {
        "titre": "Fiche SBC journalière : synthèse par aire de santé (contacts, décès, réticences)",
        "cle": ["date", "aire_sante"],
        "colonnes": [
            ("date", "date", "Date", True, ""),
            ("ip", "texte", "Partenaire", True, "liste:ip"),
            ("province", "texte", "Province", True, "liste:provinces"),
            ("zone_sante", "texte", "Zone de santé", True, ""),
            ("aire_sante", "texte", "Aire de santé", True, ""),
            ("reco_attendus", "entier", "RECO et ACSA attendus", True, ""),
            ("reco_ayant_travaille", "entier", "RECO et ACSA ayant travaillé", True, ""),
            ("contacts_listes_h", "entier", "Contacts listés ce jour : hommes", False, ""),
            ("contacts_listes_f", "entier", "Contacts listés ce jour : femmes", False, ""),
            ("contacts_sous_suivi_h", "entier", "Contacts sous suivi dans l'AS : hommes", False, ""),
            ("contacts_sous_suivi_f", "entier", "Contacts sous suivi dans l'AS : femmes", False, ""),
            ("contacts_vus_h", "entier", "Contacts vus ce jour : hommes", False, ""),
            ("contacts_vus_f", "entier", "Contacts vus ce jour : femmes", False, ""),
            ("contacts_symptomatiques", "entier", "Contacts avec au moins un signe", False, ""),
            ("contacts_symptomatiques_referes_24h", "entier",
             "NOUVEAU : contacts symptomatiques référés en 24 h", False, ""),
            ("contacts_fin_suivi_21j", "entier", "NOUVEAU : contacts ayant terminé les 21 jours",
             False, ""),
            ("contacts_perdus_de_vue", "entier", "NOUVEAU : contacts perdus de vue", False, ""),
            ("contacts_sensibilises_vaccination", "entier", "Contacts sensibilisés à la vaccination",
             False, ""),
            ("contacts_vaccines", "entier", "Contacts vaccinés", False, ""),
            ("deces_communautaires_notifies", "entier", "NOUVEAU : décès communautaires notifiés",
             False, ""),
            ("deces_communautaires_swabes", "entier", "Décès communautaires prélevés (swab)",
             False, ""),
            ("deces_communautaires_eds", "entier", "Décès communautaires avec EDS", False, ""),
            ("alertes_verifiees_24h", "entier", "NOUVEAU : alertes vérifiées par l'ESS/EIR en 24 h",
             False, ""),
            ("reticences_enregistrees", "entier", "Réticences enregistrées", False, ""),
            ("reticences_levees_meme_jour", "entier", "Réticences levées le jour même", False, ""),
            ("reticences_levees_48h", "entier", "NOUVEAU : réticences levées en 48 h", False, ""),
            ("rumeurs_enregistrees", "entier", "NOUVEAU : rumeurs et préoccupations enregistrées",
             False, ""),
            ("rumeurs_traitees_48h", "entier", "NOUVEAU : rumeurs traitées en 48 h", False, ""),
            ("alertes_sante_animale", "entier", "NOUVEAU : alertes de santé animale (One Health)",
             False, ""),
            ("feedback_majeur", "texte", "Feedbacks majeurs (un par ligne)", False, ""),
            ("visa_as", "texte", "Visa de l'AS (O/N)", False, "liste:oui_non"),
            ("visa_bcz", "texte", "Visa du BCZ (O/N)", False, "liste:oui_non"),
        ],
    },
    "reunions": {
        "titre": "Réunions de coordination",
        "cle": ["id_reunion"],
        "colonnes": [
            ("id_reunion", "texte", "Identifiant (AAAAMMJJ-niveau-lieu)", True, ""),
            ("date", "date", "Date", True, ""),
            ("niveau", "texte", "National / provincial / ZS / pilier / partenaires", True,
             "liste:niveau_reunion"),
            ("lieu", "texte", "Lieu", True, ""),
            ("objet", "texte", "Objet", True, ""),
            ("duree_min", "entier", "Durée en minutes", False, ""),
            ("nb_participants", "entier", "Nombre de participants", False, ""),
            ("representant_cgpp", "texte", "Représentant CGPP", True, ""),
            ("fichier_rapport", "texte", "Compte rendu", False, ""),
            ("liste_presence_complete", "texte", "Liste de présence complète (O/N)", False,
             "liste:oui_non"),
        ],
    },
    "points_action": {
        "titre": "Points d'action issus des réunions",
        "cle": ["id_action"],
        "colonnes": [
            ("id_action", "texte", "Identifiant", True, ""),
            ("id_reunion", "texte", "Réunion d'origine", True, ""),
            ("action", "texte", "Action à mener", True, ""),
            ("concerne_cgpp", "texte", "Implique CGPP (O/N)", True, "liste:oui_non"),
            ("responsable", "texte", "Responsable", False, ""),
            ("echeance", "date", "Échéance", False, ""),
            ("statut", "texte", "À faire / En cours / Réalisé / Abandonné", True,
             "liste:statut_action"),
            ("date_statut", "date", "Date du dernier statut", False, ""),
            ("commentaire", "texte", "Remarque", False, ""),
        ],
    },
    "autres_epidemies": {
        "titre": "Autres épidémies dans les provinces d'intervention (choléra, mpox, rougeole) et alertes des RECO CGPP",
        "cle": ["annee", "semaine_epi", "province", "maladie"],
        "colonnes": [
            ("annee", "entier", "Année", True, ""),
            ("semaine_epi", "entier", "Semaine épidémiologique", True, ""),
            ("province", "texte", "Province", True, "liste:provinces"),
            ("maladie", "texte", "Maladie", True, "liste:maladies"),
            ("zs_notifiantes", "entier", "ZS ayant notifié des cas", False, ""),
            ("zs_total", "entier", "ZS de la province", False, ""),
            ("cas", "entier", "Cas notifiés dans la semaine", False, ""),
            ("deces", "entier", "Décès notifiés dans la semaine", False, ""),
            ("zs_en_epidemie", "texte", "ZS en épidémie (liste)", False, ""),
            ("zs_cgpp_concernees", "texte", "ZS du dispositif CGPP concernées (cas ou épidémie)", False, ""),
            ("cas_suspects_referes_cgpp", "entier", "Cas suspects alertés et référés par les RECO CGPP", False, ""),
            ("detail_references", "texte", "Détail des références par ZS", False, ""),
            ("source", "texte", "Document d'origine", True, ""),
        ],
    },
    "sitrep_notes": {
        "titre": "Notes de rédaction du SitRep CGPP (faits marquants, défis, actions) saisies par l'équipe",
        "cle": ["date_rapport", "section", "ordre"],
        "colonnes": [
            ("date_rapport", "date", "Date du SitRep concerné", True, ""),
            ("section", "texte", "Faits marquants / Activités / Sécurité / Défis / Actions / Commentaire", True,
             "liste:section_note"),
            ("ordre", "entier", "Ordre d'affichage", True, ""),
            ("texte", "texte", "Texte (une idée par ligne, phrase complète)", True, ""),
            ("auteur", "texte", "Auteur", False, ""),
        ],
    },
    "qc_suivi": {
        "titre": "Suivi des anomalies de qualité des données",
        "cle": ["id_anomalie"],
        "colonnes": [
            ("id_anomalie", "texte", "Identifiant (règle + objet)", True, ""),
            ("date_detection", "date", "Date de détection", True, ""),
            ("regle", "texte", "Règle de contrôle", True, ""),
            ("source", "texte", "Source concernée", True, ""),
            ("objet", "texte", "Élément concerné", False, ""),
            ("description", "texte", "Description", True, ""),
            ("gravite", "texte", "Critique / Majeure / Mineure", True, "liste:gravite"),
            ("responsable", "texte", "Responsable de la correction", False, ""),
            ("statut", "texte", "Ouverte / Corrigée / Acceptée", True, "liste:statut_anomalie"),
            ("date_resolution", "date", "Date de résolution", False, ""),
        ],
    },
}

LISTES = {
    "provinces": PROVINCES,
    "ip": ["WV", "ADRA", "IMC", "IRC", "CGPP"],
    "statut_cgpp": ["intervention", "préparation", "hors zone CGPP"],
    "piliers": ["COORD", "SURV", "CREC", "PCI", "LOG", "MEL"],
    "type_indicateur": ["Nombre", "Pourcentage", "Délai", "Ratio"],
    "statut_indicateur": ["existant", "à corriger", "nouveau"],
    "statut_donnees": ["OK", "ND"],
    "types_site": TYPES_SITE,
    "oui_non": ["O", "N"],
    "niveau_reunion": ["National", "Provincial", "Zone de santé", "Pilier", "Partenaires", "Interne CGPP"],
    "statut_action": ["À faire", "En cours", "Réalisé", "Abandonné"],
    "gravite": ["Critique", "Majeure", "Mineure"],
    "statut_anomalie": ["Ouverte", "Corrigée", "Acceptée"],
    "section_note": ["Faits marquants", "Activités", "Sécurité", "Défis", "Actions", "Commentaire"],
    "maladies": ["Choléra", "Mpox", "Rougeole", "Autre"],
}


def colonnes(feuille: str) -> list[str]:
    return [c[0] for c in SCHEMA[feuille]["colonnes"]]


def table_vide(feuille: str) -> pd.DataFrame:
    return pd.DataFrame(columns=colonnes(feuille))


# --------------------------------------------------------------------------
# Lecture / écriture de la base Excel
# --------------------------------------------------------------------------

def lire_table(chemin_base: Path, feuille: str) -> pd.DataFrame:
    try:
        df = pd.read_excel(chemin_base, sheet_name=feuille)
    except (FileNotFoundError, ValueError):
        return table_vide(feuille)
    for c in colonnes(feuille):
        if c not in df.columns:
            df[c] = None
    return df[colonnes(feuille)]


def ecrire_tables(chemin_base: Path, tables: dict[str, pd.DataFrame], canevas: Path | None = None):
    """Remplace le contenu des feuilles indiquées en conservant en-têtes, formats et listes.
    Si la base n'existe pas, elle est créée à partir du canevas."""
    import openpyxl
    from openpyxl.utils import get_column_letter

    chemin_base = Path(chemin_base)
    if not chemin_base.exists():
        if canevas and Path(canevas).exists():
            shutil.copy(canevas, chemin_base)
            wb0 = openpyxl.load_workbook(chemin_base)
            for nom in SCHEMA:                       # retirer les lignes d'exemple du canevas
                if nom in wb0.sheetnames and not nom.startswith("ref_") and wb0[nom].max_row > 1:
                    wb0[nom].delete_rows(2, wb0[nom].max_row - 1)
            wb0.save(chemin_base)
        else:
            wb = openpyxl.Workbook()
            wb.active.title = "LISEZMOI"
            wb.active["A1"] = "Base de données CGPP MVE-17 (créée sans canevas)"
            wb.save(chemin_base)
    wb = openpyxl.load_workbook(chemin_base)
    for feuille, df in tables.items():
        cols = colonnes(feuille)
        if feuille not in wb.sheetnames:
            ws = wb.create_sheet(feuille)
            for j, c in enumerate(cols, start=1):
                ws.cell(row=1, column=j, value=c)
        ws = wb[feuille]
        entetes = [ws.cell(row=1, column=j).value for j in range(1, ws.max_column + 1)]
        if [e for e in entetes if e] != cols:
            for j, c in enumerate(cols, start=1):
                ws.cell(row=1, column=j, value=c)
        if ws.max_row > 1:
            ws.delete_rows(2, ws.max_row - 1)
        df = df.reindex(columns=cols)
        types = [c[1] for c in SCHEMA[feuille]["colonnes"]]
        for i, ligne in enumerate(df.itertuples(index=False), start=2):
            for j, v in enumerate(ligne, start=1):
                type_col = types[j - 1]
                if type_col.startswith("formule:"):
                    cell = ws.cell(row=i, column=j, value=type_col[len("formule:"):].replace("{r}", str(i)))
                    cell.number_format = "0.0"
                    continue
                if v is None or (isinstance(v, float) and pd.isna(v)):
                    continue
                if isinstance(v, pd.Timestamp):
                    v = v.to_pydatetime()
                cell = ws.cell(row=i, column=j, value=v)
                if type_col == "date":
                    cell.number_format = "DD/MM/YYYY"
                elif type_col == "decimal":
                    cell.number_format = "0.0"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{max(2, len(df) + 1)}"
    wb.calculation.fullCalcOnLoad = True   # Excel recalcule les formules à l'ouverture
    wb.save(chemin_base)


def valeur_indicateur(df: pd.DataFrame) -> pd.Series:
    """Valeur d'une ligne de saisie : numérateur / dénominateur x 100, ou numérateur seul."""
    num = pd.to_numeric(df["numerateur"], errors="coerce")
    den = pd.to_numeric(df["denominateur"], errors="coerce")
    return num.where(den.isna() | (den == 0), 100 * num / den)


def cible_numerique(texte):
    """'2 000' -> 2000 ; '≥ 90 %' -> 90 ; '100 %' -> 100 ; 'Suivi' ou texte composé -> None."""
    if texte is None or (isinstance(texte, float) and pd.isna(texte)):
        return None
    t = str(texte).strip()
    m = re.fullmatch(r"[≥>]?\s*([\d\s  ]+(?:[.,]\d+)?)\s*%?", t)
    return nombre_fr(m.group(1)) if m else None


def fusionner(existant: pd.DataFrame, nouveau: pd.DataFrame, cle_cols: list[str]) -> pd.DataFrame:
    """Remplace dans `existant` les lignes ayant la même clé que `nouveau` (upsert)."""
    if existant is None or existant.empty:
        return nouveau.reset_index(drop=True)
    if nouveau.empty:
        return existant

    def norm(v):
        if v is None or (isinstance(v, float) and pd.isna(v)) or v is pd.NaT:
            return ""
        if hasattr(v, "isoformat"):            # date, datetime, Timestamp -> AAAA-MM-JJ
            return pd.Timestamp(v).date().isoformat()
        if isinstance(v, float) and v.is_integer():
            return str(int(v))
        return str(v).strip()

    def cles(df):
        return df[cle_cols].apply(lambda col: col.map(norm)).agg("|".join, axis=1)

    garde = existant.loc[~cles(existant).isin(set(cles(nouveau)))]
    morceaux = [x for x in (garde, nouveau) if not x.empty]   # pas de concaténation de tables vides (pandas 2.x)
    if not morceaux:
        return nouveau.copy()
    if len(morceaux) == 1:
        return morceaux[0].reset_index(drop=True)
    with warnings.catch_warnings():   # colonnes entièrement vides : avertissement de pandas 2.x sans effet ici
        warnings.simplefilter("ignore", FutureWarning)
        return pd.concat(morceaux, ignore_index=True)
