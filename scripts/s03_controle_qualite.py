"""
s03_controle_qualite.py - Contrôle de la qualité des données (règles automatisées).

Dimensions contrôlées : exhaustivité, cohérence interne (sommes, recalcul des pourcentages),
cohérence externe (entre documents), cohérence temporelle (entre deux SitReps), conformité au
cadre d'indicateurs, format et protection des données.

Entrées : base Excel (s01, s02), fichiers intermédiaires, documents sources (réunions, plans,
          matrice de Dr Barclay, cadre global, fiche SBC), observations manuelles.
Sorties : sorties/qc/rapport_qc_AAAA-MM-JJ.xlsx (synthèse, anomalies, tous les contrôles,
          pourcentages recalculés) et mise à jour de la feuille qc_suivi de la base.

Usage   : python scripts/s03_controle_qualite.py
"""
from __future__ import annotations

import re
import sys
from datetime import date
from pathlib import Path

import docx
import openpyxl
import pandas as pd

sys.path.insert(0, str(Path(__file__).resolve().parent))
from commun import (SCHEMA, charger_config, cle, ecrire_tables, fusionner, journal, lire_table,  # noqa: E402
                    nombre_fr, vers_docx)
from s01_extraire_sitreps import verifier_pourcentages  # noqa: E402

LOG = journal("qc")
R: list[dict] = []   # résultats


def res(id_regle, dimension, source, objet, description, attendu=None, observe=None, statut="OK",
        gravite="", recommandation=""):
    R.append(dict(id_regle=id_regle, dimension=dimension, source=source, objet=objet, description=description,
                  attendu=attendu, observe=observe, statut=statut, gravite=gravite if statut != "OK" else "",
                  recommandation=recommandation if statut != "OK" else ""))


def egal(a, b, tol=0.0):
    if a is None or b is None or pd.isna(a) or pd.isna(b):
        return None
    return abs(float(a) - float(b)) <= tol


def f(x):
    """Affichage d'un nombre à la française."""
    if x is None or (isinstance(x, float) and pd.isna(x)):
        return "ND"
    if float(x).is_integer():
        return f"{int(x):,}".replace(",", " ")
    return f"{x:.1f}".replace(".", ",")


# ==========================================================================
# A. SitReps nationaux
# ==========================================================================

def controles_sitreps(cfg, base):
    meta = lire_table(base, "sitrep_meta")
    prov = lire_table(base, "epi_province")
    zs = lire_table(base, "epi_zs")
    alertes = lire_table(base, "alertes_province")
    contacts = lire_table(base, "contacts_province")
    labo = lire_table(base, "labo_province")
    pec = lire_table(base, "pec_province")
    poe = lire_table(base, "poe_province")
    textes = lire_table(base, "sitrep_texte")
    ctrl = {c["sitrep_no"]: c for c in pd.read_pickle(cfg["chemins"]["intermediaire"] / "controles_sitreps.pkl")}

    nos = sorted(meta["sitrep_no"].astype(int))
    for no in nos:
        src = f"SitRep {no}"
        m = meta[meta.sitrep_no == no].iloc[0]
        p = prov[prov.sitrep_no == no]
        z = zs[zs.sitrep_no == no]
        zv = z[~z.est_a_ventiler.fillna(False).astype(bool)]
        tot2 = ctrl.get(no, {}).get("totaux_t2", {})

        # SR-01 somme des ZS = province
        for r in p.itertuples():
            zp = z[z.province == r.province]
            for champ, lib in [("cas_cumul", "cas"), ("deces_cumul", "décès")]:
                s = zp[champ].fillna(0).sum()
                ok = egal(s, getattr(r, champ))
                res("SR-01", "Cohérence interne", src, f"{r.province} : {lib} cumulés",
                    "Somme des ZS = total de la province", f(getattr(r, champ)), f(s),
                    "OK" if ok else "ÉCHEC", "Critique", "Corriger le tableau 2 avant publication.")
        # SR-02 somme des provinces = total national = encadré
        for champ, lib in [("cas_cumul", "cas"), ("deces_cumul", "décès")]:
            s = p[champ].sum()
            ok = egal(s, m[champ])
            res("SR-02", "Cohérence interne", src, f"National : {lib} cumulés",
                "Somme des provinces = chiffre de l'encadré", f(m[champ]), f(s), "OK" if ok else "ÉCHEC", "Critique",
                "Aligner l'encadré et le tableau 1.")
        # SR-03 / SR-04 létalité recalculée
        for r in p.itertuples():
            calc = round(100 * r.deces_cumul / r.cas_cumul, 1) if r.cas_cumul else None
            ok = egal(calc, r.letalite_publiee, 0.1)
            res("SR-03", "Cohérence interne", src, f"{r.province} : létalité", "Létalité publiée = décès / cas",
                f(r.letalite_publiee), f(calc), "OK" if ok else "ÉCHEC", "Majeure", "Recalculer la létalité.")
        n_ko = 0
        for r in zv.itertuples():
            if not r.cas_cumul:
                continue
            calc = round(100 * r.deces_cumul / r.cas_cumul, 1)
            if not egal(calc, r.letalite_publiee, 0.11):
                n_ko += 1
                res("SR-04", "Cohérence interne", src, f"{r.zone_sante} : létalité", "Létalité ZS recalculée",
                    f(r.letalite_publiee), f(calc), "ÉCHEC", "Mineure", "Recalculer la létalité de la ZS.")
        if n_ko == 0:
            res("SR-04", "Cohérence interne", src, f"{len(zv)} ZS : létalité", "Létalité ZS recalculée",
                "écart ≤ 0,1 point", "toutes conformes")
        # SR-05 nouveaux cas : ZS = province = national = faits saillants
        for r in p.itertuples():
            s = z[z.province == r.province]["nouveaux_cas_24h"].fillna(0).sum()
            ok = egal(s, r.nouveaux_cas_24h)
            res("SR-05", "Cohérence interne", src, f"{r.province} : nouveaux cas 24 h", "Somme des ZS = province",
                f(r.nouveaux_cas_24h), f(s), "OK" if ok else "ÉCHEC", "Critique")
        s = p["nouveaux_cas_24h"].sum()
        ok = egal(s, m.nouveaux_cas_24h)
        res("SR-05", "Cohérence interne", src, "National : nouveaux cas 24 h", "Tableau 1 = faits saillants",
            f(m.nouveaux_cas_24h), f(s), "OK" if ok else "ÉCHEC", "Critique")
        # SR-06 décès du jour
        t = tot2.get("Total", {})
        if t:
            ok = egal((t.get("deces_communautaires_24h") or 0) + (t.get("deces_cte_24h") or 0), t.get("deces_total_24h"))
            res("SR-06", "Cohérence interne", src, "Décès du jour", "Communautaires + intra-CTE = total",
                f(t.get("deces_total_24h")), f((t.get("deces_communautaires_24h") or 0) + (t.get("deces_cte_24h") or 0)),
                "OK" if ok else "ÉCHEC", "Majeure")
            ok = egal(t.get("deces_communautaires_24h"), m.deces_communautaires_24h) and egal(t.get("deces_cte_24h"),
                                                                                             m.deces_cte_24h)
            res("SR-06", "Cohérence interne", src, "Décès du jour", "Tableau 2 = faits saillants",
                f"{f(m.deces_communautaires_24h)} + {f(m.deces_cte_24h)}",
                f"{f(t.get('deces_communautaires_24h'))} + {f(t.get('deces_cte_24h'))}", "OK" if ok else "ÉCHEC", "Majeure")
        # SR-07 ZS touchées
        for r in p.itertuples():
            n = len(zv[zv.province == r.province])
            ok = egal(n, r.zs_touchees)
            res("SR-07", "Cohérence interne", src, f"{r.province} : ZS touchées", "Lignes ZS du tableau 2 = tableau 1",
                f(r.zs_touchees), f(n), "OK" if ok else "ÉCHEC", "Majeure")
        # SR-08 alertes
        a = alertes[alertes.sitrep_no == no]
        ok_a = a[a.statut_donnees == "OK"]
        for r in ok_a.itertuples():
            v = egal(r.recues_vivants + r.recues_deces, r.recues_total) and \
                egal(r.validees_vivants + r.validees_deces + r.invalidees_vivants + r.invalidees_deces, r.recues_total)
            res("SR-08", "Cohérence interne", src, f"{r.province} : alertes",
                "Vivants + décès = total ; validées + invalidées = total", f(r.recues_total),
                f(r.validees_vivants + r.validees_deces + r.invalidees_vivants + r.invalidees_deces),
                "OK" if v else "ÉCHEC", "Majeure")
        nd = sorted(a[a.statut_donnees == "ND"]["province"])
        if nd:
            res("SR-08", "Exhaustivité", src, "Alertes", "Toutes les provinces touchées rapportent",
                "7 provinces", f"ND : {', '.join(nd)}", "ÉCHEC", "Critique",
                "Afficher « provinces rapportantes x/y » à côté de chaque total et ne pas comparer au jour précédent.")
        txt = ctrl.get(no, {}).get("alertes_texte", {})
        tot_alertes = ok_a["recues_total"].sum()
        verif_table = 100.0 if len(ok_a) else None
        if txt.get("pct_verifiees_texte") is not None:
            ok = egal(txt["pct_verifiees_texte"], verif_table, 0.1)
            res("SR-08", "Cohérence interne", src, "Alertes vérifiées (texte)", "Texte = tableau 3",
                f"{f(verif_table)} %", f"{f(txt['pct_verifiees_texte'])} %", "OK" if ok else "ÉCHEC", "Majeure",
                "Le tableau ne contient aucune alerte non vérifiée : corriger le pourcentage du texte.")
        if txt.get("pct_validees_texte") is not None and tot_alertes:
            calc = round(100 * txt["validees_texte"] / tot_alertes, 1)
            ok = egal(calc, txt["pct_validees_texte"], 0.15)
            res("SR-08", "Cohérence interne", src, "Alertes validées (texte)", "Pourcentage recalculé",
                f"{f(calc)} %", f"{f(txt['pct_validees_texte'])} %", "OK" if ok else "ÉCHEC", "Mineure")
        # SR-09 contacts
        c = contacts[contacts.sitrep_no == no]
        cok = c[c.statut_donnees == "OK"]
        tot = ctrl.get(no, {}).get("contacts_total", {})
        if tot:
            ok = egal(cok.contacts_suivis.sum(), tot.get("contacts_suivis")) and egal(cok.contacts_vus.sum(),
                                                                                       tot.get("contacts_vus"))
            res("SR-09", "Cohérence interne", src, "Contacts", "Somme des provinces = total",
                f"{f(tot.get('contacts_vus'))}/{f(tot.get('contacts_suivis'))}",
                f"{f(cok.contacts_vus.sum())}/{f(cok.contacts_suivis.sum())}", "OK" if ok else "ÉCHEC", "Majeure")
        cnd = sorted(c[c.statut_donnees == "ND"]["province"])
        if cnd:
            res("SR-09", "Exhaustivité", src, "Taux de suivi des contacts (encadré)",
                "Le taux national inclut toutes les provinces", "toutes les provinces",
                f"sans {', '.join(cnd)}", "ÉCHEC", "Critique",
                "Afficher le taux sur les provinces rapportantes et le signaler dans l'encadré ; "
                "comparer au jour précédent sur le même périmètre.")
        # SR-10 laboratoire
        lb = labo[labo.sitrep_no == no]
        if len(lb):
            ok = egal(lb.echantillons.sum(), m.echantillons_24h)
            res("SR-10", "Cohérence interne", src, "Échantillons analysés", "Somme des provinces = faits saillants",
                f(m.echantillons_24h), f(lb.echantillons.sum()), "OK" if ok else "ÉCHEC", "Majeure")
            ok = egal(lb.positifs.sum(), m.nouveaux_cas_24h)
            res("SR-10", "Cohérence externe", src, "Résultats positifs", "Positifs du laboratoire = nouveaux cas",
                f(m.nouveaux_cas_24h), f(lb.positifs.sum()), "OK" if ok else "À VÉRIFIER", "Mineure",
                "Préciser si des positifs proviennent d'autres laboratoires ou de reclassements.")
            for r in lb.itertuples():
                if r.positivite_publiee is not None and not pd.isna(r.positivite_publiee) and r.echantillons:
                    calc = round(100 * r.positifs / r.echantillons, 1)
                    ok = egal(calc, r.positivite_publiee, 0.15)
                    res("SR-10", "Cohérence interne", src, f"{r.province} : positivité", "Positivité recalculée",
                        f"{f(r.positivite_publiee)} %", f"{f(calc)} %", "OK" if ok else "ÉCHEC", "Mineure")
        # SR-12 patients en isolement / SR-13 guéris
        pc = pec[pec.sitrep_no == no]
        if len(pc):
            s = pc.patients_isolement.sum()
            ok = egal(s, m.patients_isolement)
            res("SR-12", "Cohérence interne", src, "Patients en isolement", "Encadré = somme des provinces",
                f(m.patients_isolement), f(s), "OK" if ok else "ÉCHEC", "Majeure",
                "Vérifier la Tshopo (« 15 patients (6 confirmés et 4 suspects) »)." if no == 146 else "")
            g = pc.guerisons.fillna(0).sum()
            ok = egal(g, m.gueris_24h)
            detail = "; ".join(f"{r.province} {f(r.guerisons)}" for r in pc.itertuples() if r.guerisons)
            res("SR-13", "Cohérence interne", src, "Guéris du jour", "Résumé = somme des guérisons par province",
                f(m.gueris_24h), f"{f(g)} ({detail})", "OK" if ok else "ÉCHEC", "Majeure",
                "Le résumé omet une province : recalculer le total des guéris.")
        # SR-15 décès à ventiler
        av = z[z.est_a_ventiler.fillna(False).astype(bool)]
        for r in av.itertuples():
            tot_prov = p[p.province == r.province]["deces_cumul"].sum()
            part = 100 * r.deces_cumul / tot_prov if tot_prov else 0
            res("SR-15", "Exhaustivité", src, f"{r.province} : décès à ventiler",
                "Décès attribués à une ZS (part non ventilée < 5 %)", "< 5 %",
                f"{f(r.deces_cumul)} ({part:.1f} %)".replace(".", ","), "ÉCHEC" if part >= 5 else "OK", "Majeure",
                "Ventiler les décès par ZS : les létalités par ZS sont sous-estimées tant que ce n'est pas fait.")
        # SR-19 PoE : global = somme
        pz = poe[poe.sitrep_no == no]
        for ind in ["Nombre de personnes passées aux PoE/PoC", "Nombre d'alertes notifiées"]:
            q = pz[pz.indicateur.str.startswith(ind[:20])]
            if len(q):
                glob = q[q.province == "Global"]["valeur"].sum()
                s = q[q.province != "Global"]["valeur"].sum()
                ok = egal(glob, s)
                res("SR-19", "Cohérence interne", src, f"PoE/PoC : {ind}", "Global = somme des provinces", f(glob), f(s),
                    "OK" if ok else "ÉCHEC", "Mineure")

    # SR-11 pourcentages du texte
    pcts = pd.read_csv(cfg["chemins"]["intermediaire"] / "pourcentages_texte_sitreps.csv")
    for r in pcts.itertuples():
        if r.statut == "ÉCART":
            res("SR-11", "Cohérence interne", r.reference, "Pourcentage écrit dans le texte",
                f"« {r.extrait.strip()[-110:]} »", f"{f(r.pct_recalcule)} % ({f(r.numerateur)}/{f(r.denominateur)})",
                f"{f(r.pct_publie)} % publié", "ÉCHEC", "Majeure", "Corriger le pourcentage ou la fraction.")
    for ref, g in pcts.groupby("reference"):
        n_ok = (g.statut == "OK").sum()
        res("SR-11", "Cohérence interne", ref, f"{len(g)} pourcentages recalculés", "Pourcentage = numérateur/dénominateur",
            f"{len(g)} conformes", f"{n_ok} conformes", "OK" if n_ok == len(g) else "INFO")

    # SR-20 « N patients (a confirmés et b suspects) » : a + b = N
    for r in textes.itertuples():
        for mm in re.finditer(r"(\d+)\s+patients\s+\((\d+)\s+confirmés?\s+et\s+(\d+)\s+suspects?\)", str(r.texte)):
            n, a_, b_ = map(int, mm.groups())
            res("SR-20", "Cohérence interne", f"SitRep {r.sitrep_no}", "Ventilation des patients",
                "Total = confirmés + suspects", f"{a_} + {b_} = {a_ + b_}", f"{n} patients annoncés",
                "OK" if a_ + b_ == n else "ÉCHEC", "Mineure", "Corriger le total ou la ventilation.")

    # SR-14 taux d'occupation : numérateur repris d'un SitRep précédent
    occ = []
    for no in nos:
        for t in textes[(textes.sitrep_no == no) & (textes.section == "Prise en charge")]["texte"]:
            mm = re.search(r"(\d+)\s+patients[^.]*?en isolement.*?occupation[^()]*?\((\d[\d\s ]*)/(\d[\d\s ]*)\)", t)
            mp = re.match(r"^\W*(?:En|Au|À la)\s+(\w[\w\-é]+)", t)
            if mm and mp:
                occ.append(dict(no=no, prov=mp.group(1), patients=int(mm.group(1)), num=nombre_fr(mm.group(2)),
                                den=nombre_fr(mm.group(3))))
    prec = {}
    for o in occ:
        pi = pec[(pec.sitrep_no < o["no"]) & (pec.province.str.startswith(o["prov"][:4]))]
        prev_patients = pi.sort_values("sitrep_no")["patients_isolement"].iloc[-1] if len(pi) else None
        if o["num"] == prev_patients and o["num"] != o["patients"]:
            res("SR-14", "Cohérence temporelle", f"SitRep {o['no']}", f"{o['prov']} : taux d'occupation des lits",
                "Numérateur = patients du jour", f"{o['patients']} patients",
                f"{f(o['num'])}/{f(o['den'])}, soit le chiffre du SitRep précédent", "ÉCHEC", "Majeure",
                "Valeur reprise du SitRep précédent : recalculer avec les patients du jour.")
        prec[o["prov"]] = o

    # B. comparaison entre SitReps successifs
    for i in range(1, len(nos)):
        a_no, b_no = nos[i - 1], nos[i]
        za = zs[(zs.sitrep_no == a_no) & ~zs.est_a_ventiler.fillna(False).astype(bool)].set_index(["province", "zone_sante"])
        zb = zs[(zs.sitrep_no == b_no) & ~zs.est_a_ventiler.fillna(False).astype(bool)].set_index(["province", "zone_sante"])
        baisses, incoherents = [], []
        for k, r in zb.iterrows():
            if k in za.index:
                delta = (r.cas_cumul or 0) - (za.loc[k].cas_cumul or 0)
                if delta < 0:
                    baisses.append(f"{k[1]} ({f(delta)})")
                if (r.nouveaux_cas_24h or 0) > delta:
                    incoherents.append(f"{k[1]} (hausse {f(delta)}, nouveaux {f(r.nouveaux_cas_24h)})")
        res("TS-01", "Cohérence temporelle", f"SitRep {a_no} → {b_no}", "Cumuls par ZS", "Aucun cumul ne diminue",
            "0 baisse", f"{len(baisses)} baisse(s)" + (f" : {', '.join(baisses)}" if baisses else ""),
            "OK" if not baisses else "ÉCHEC", "Majeure", "Documenter les reclassements de cas.")
        res("TS-02", "Cohérence temporelle", f"SitRep {a_no} → {b_no}", "Hausse des cumuls vs nouveaux cas du jour",
            "Hausse du cumul ≥ nouveaux cas du dernier jour", "toutes les ZS",
            f"{len(incoherents)} ZS incohérente(s)" + (f" : {', '.join(incoherents)}" if incoherents else ""),
            "OK" if not incoherents else "ÉCHEC", "Majeure")
        nouvelles = sorted(set(zb.index) - set(za.index))
        res("TS-03", "Cohérence temporelle", f"SitRep {a_no} → {b_no}", "Nouvelles ZS touchées", "Information",
            "", ", ".join(f"{k[1]} ({k[0]})" for k in nouvelles) or "aucune", "INFO" if nouvelles else "OK")
        ma, mb = meta[meta.sitrep_no == a_no].iloc[0], meta[meta.sitrep_no == b_no].iloc[0]
        if (mb.zs_touchees if "zs_touchees" in mb else None) is None and nouvelles and \
                ma.aires_sante_touchees == mb.aires_sante_touchees:
            res("SR-16", "Cohérence temporelle", f"SitRep {b_no}", "Aires de santé touchées",
                "Le nombre d'AS augmente quand une ZS nouvelle est touchée", f"> {f(ma.aires_sante_touchees)}",
                f(mb.aires_sante_touchees), "ÉCHEC", "Mineure", "Mettre à jour le compteur d'aires de santé.")
        # comparabilité de l'encadré « taux de suivi des contacts »
        cb = contacts[(contacts.sitrep_no == b_no)]
        if (cb.statut_donnees == "ND").any():
            res("TS-04", "Cohérence temporelle", f"SitRep {a_no} → {b_no}", "Évolution du taux de suivi des contacts",
                "Comparaison sur le même périmètre", f"{f(ma.taux_suivi_contacts)} % (7 provinces)",
                f"{f(mb.taux_suivi_contacts)} % (sans {', '.join(cb[cb.statut_donnees == 'ND'].province)})",
                "ÉCHEC", "Critique", "Ne pas présenter la variation comme une évolution réelle.")


# ==========================================================================
# C. Comptes rendus de coordination
# ==========================================================================

def controles_reunions(cfg, base):
    dossier = cfg["chemins"]["entree_reunions"]
    for fch in sorted(dossier.iterdir()):
        if fch.suffix.lower() not in (".doc", ".docx"):
            continue
        d = docx.Document(vers_docx(fch, cfg["chemins"]["intermediaire"], cfg.get("libreoffice", "soffice")))
        texte = "\n".join(p.text for p in d.paragraphs)
        for t in d.tables:
            for row in t.rows:
                texte += "\n" + " | ".join(c.text for c in row.cells)
        for r in verifier_pourcentages(texte, fch.name, fch.stem, cfg["seuils"]["tolerance_pourcentage"]):
            res("CR-01", "Cohérence interne", fch.stem, "Pourcentage écrit dans le compte rendu",
                f"« {r['extrait'].strip()[-90:]} »", f"{f(r['pct_recalcule'])} %", f"{f(r['pct_publie'])} % publié",
                "OK" if r["statut"] == "OK" else "ÉCHEC", "Majeure" if r["note"] else "Mineure",
                ((r["note"] + " ; ") if r["note"] else "") + "Corriger avant diffusion.")
    hist = lire_table(base, "historique")
    hist["date_reference"] = pd.to_datetime(hist["date_reference"])
    it = hist[(hist.territoire == "Ituri") & (hist.indicateur.isin(["cas_cumul", "nouveaux_cas_24h"]))]
    try:
        c4 = it[(it.date_reference == "2026-10-04") & (it.indicateur == "cas_cumul")].valeur.iloc[0]
        c5 = it[(it.date_reference == "2026-10-05") & (it.indicateur == "cas_cumul")].valeur.iloc[0]
        n5 = it[(it.date_reference == "2026-10-05") & (it.indicateur == "nouveaux_cas_24h")].valeur.iloc[0]
        res("CR-02", "Cohérence externe", "CR 05/10 et 06/10", "Ituri : cumul des cas",
            "Cumul J = cumul J-1 + nouveaux cas", f"{f(c4)} + {f(n5)} = {f(c4 + n5)}", f"{f(c5)}",
            "ÉCHEC" if c4 + n5 != c5 else "OK", "Majeure",
            "Le cumul du 4/10 est probablement 6 392 (et non 6 292).")
        prov = lire_table(base, "epi_province")
        s146 = prov[(prov.sitrep_no == 146) & (prov.province == "Ituri")].iloc[0]
        implicite = s146.cas_cumul - s146.nouveaux_cas_24h - c5
        res("CR-02", "Cohérence externe", "CR 06/10 vs SitRep 146", "Ituri : cumul du 5/10 au 7/10",
            "Hausse plausible le 6/10 (25 à 50 cas/jour)", "25 à 50",
            f"{f(implicite)} cas implicites le 6/10 ({f(c5)} → {f(s146.cas_cumul)} avec {f(s146.nouveaux_cas_24h)} le 7/10)",
            "À VÉRIFIER" if implicite > 60 else "OK", "Majeure",
            "Réconcilier avec le SitRep 145 ; aligner les chiffres présentés en réunion sur le SitRep publié.")
    except IndexError:
        pass
    prov = lire_table(base, "epi_province")
    p146 = prov[prov.sitrep_no == 146].set_index("province")
    for territoire, (num, den) in {"Bas-Uélé": (3, 13), "Haut-Uélé": (6, 13)}.items():
        if territoire in p146.index:
            r = p146.loc[territoire]
            ok = num == r.zs_touchees and den == r.zs_total
            res("CR-03", "Cohérence externe", "CR coordination nationale 05/10", f"{territoire} : ZS touchées",
                "Même fraction que le SitRep", f"{f(r.zs_touchees)}/{f(r.zs_total)} (SitRep 146)", f"{num}/{den}",
                "OK" if ok else "ÉCHEC", "Mineure", "Utiliser le dénominateur officiel du SitRep.")
    pa = lire_table(base, "points_action")
    sans = pa[pa.responsable.isna() | pa.echeance.isna()]
    res("CR-04", "Conformité", "Comptes rendus de coordination", "Points d'action",
        "Chaque point d'action a un responsable et une échéance", f"{len(pa)} complets",
        f"{len(sans)} sur {len(pa)} sans responsable ni échéance", "OK" if sans.empty else "ÉCHEC", "Majeure",
        "Ajouter au gabarit un tableau action / responsable / échéance / statut (COORD-12).")


# ==========================================================================
# D. Classeur hebdomadaire et cadre global
# ==========================================================================

def controles_classeur(cfg, base):
    comp = pd.read_csv(cfg["chemins"]["intermediaire"] / "completude_classeur.csv")
    taux = comp.renseigne.mean() * 100 if len(comp) else 0
    smax = comp.semaine.max()
    res("IND-01", "Exhaustivité", "Classeur hebdomadaire CGPP", f"Cellules IP x ZS x semaine (S36 à S{smax})",
        "≥ 80 % des cellules renseignées", "≥ 80 %",
        f"{taux:.1f} % ({int(comp.renseigne.sum())}/{len(comp)})".replace(".", ","), "OK" if taux >= 80 else "ÉCHEC",
        "Critique", "Relancer les points focaux ; la mise à jour DOS du lundi part incomplète.")
    for (ip,), g in comp.groupby(["ip"]):
        t = g.renseigne.mean() * 100
        semaines = sorted(g[g.renseigne].semaine.unique())
        res("IND-01", "Exhaustivité", "Classeur hebdomadaire CGPP", f"{ip} : complétude", "≥ 80 %", "≥ 80 %",
            f"{t:.1f} % ; semaines saisies : {', '.join(map(str, semaines)) or 'aucune'}".replace(".", ","),
            "OK" if t >= 80 else "ÉCHEC", "Critique")
    fo = pd.read_csv(cfg["chemins"]["intermediaire"] / "formules_classeur.csv")
    pct = fo[fo.est_pourcentage & fo.somme_semaines]
    res("IND-02", "Conformité", "Classeur hebdomadaire CGPP (feuille compilée)", "Indicateurs en pourcentage",
        "Les % se calculent à partir de numérateurs et dénominateurs", "0 % additionné",
        f"{len(pct)} indicateurs en % additionnés entre semaines et entre IP", "ÉCHEC" if len(pct) else "OK", "Critique",
        "Saisir numérateur et dénominateur ; calculer le % après agrégation (feuille saisie_hebdo).")
    wb = openpyxl.load_workbook(cfg["chemins"]["classeur_indicateurs"], data_only=True)
    ws = next(wb[n] for n in wb.sheetnames if "compil" in n.lower())
    cibles = sum(1 for r in range(4, ws.max_row + 1) if ws.cell(r, 3).value and not ws.cell(r, 4).value)
    lignes = sum(1 for r in range(4, ws.max_row + 1) if ws.cell(r, 3).value)
    res("IND-03", "Exhaustivité", "Classeur hebdomadaire CGPP", "Cibles", "Une cible par indicateur", f"{lignes}",
        f"{lignes - cibles} cible(s) saisie(s) ; #DIV/0! sur {cibles} lignes", "ÉCHEC" if cibles else "OK", "Majeure",
        "Reporter les cibles de la matrice maîtresse (colonne cible_dos).")
    ips = {n.split("-")[-1] for n in wb.sheetnames if n.startswith("Indicator Framework-")}
    manquants = sorted(set(cfg["partenaires"]) - ips)
    res("IND-04", "Exhaustivité", "Classeur hebdomadaire CGPP", "Partenaires", "Tous les IP des plans de travail ont une feuille",
        ", ".join(sorted(cfg["partenaires"])), "absent : " + ", ".join(manquants) if manquants else "complet",
        "ÉCHEC" if manquants else "OK", "Majeure", "Ajouter IRC et ses ZS.")
    # IND-05 : cas confirmés déclarés vs SitRep
    s = lire_table(base, "saisie_hebdo")
    declare = pd.to_numeric(s[s.id_indicateur == "SURV-17"]["numerateur"], errors="coerce").sum()
    zs = lire_table(base, "epi_zs")
    dern = zs.sitrep_no.max()
    cg = zs[(zs.sitrep_no == dern) & zs.zone_sante.isin(cfg["zones_cgpp"].keys())]["cas_cumul"].sum()
    res("IND-05", "Cohérence externe", "Classeur hebdomadaire CGPP vs SitRep", "SURV-17 cas confirmés dans les ZS CGPP",
        f"Valeur du SitRep {dern}", f(cg), f"{f(declare)} déclaré par les IP", "ÉCHEC" if declare != cg else "OK",
        "Critique", "Calculer SURV-17 automatiquement depuis la feuille epi_zs.")
    base_agg = s[(s.periode == "S20-S35")]
    if len(base_agg):
        res("IND-06", "Format", "Classeur hebdomadaire CGPP", "Période S20 à S35",
            "Données ventilées par ZS et par semaine", "par ZS et semaine", "un seul total par IP", "ÉCHEC", "Mineure",
            "Ventiler rétroactivement si les registres le permettent.")
    # IND-07 / IND-08 cadre global
    csv = pd.read_csv(cfg["chemins"]["cadre_global_csv"], skiprows=2)
    der = csv.dropna(subset=["Indicator"]).iloc[-1]
    if "IDSR" in str(der["Indicator"]) and "supervision" in str(der["Definition"]).lower():
        res("IND-07", "Conformité", "Cadre global (CSV)", der["Indicator"].strip(), "Intitulé et définition concordent",
            "définition IDSR", "définition des visites de supervision", "ÉCHEC", "Majeure",
            "Rétablir : numérateur = alertes CGPP retrouvées dans la SIMR ; dénominateur = alertes CGPP détectées.")
    wbd = openpyxl.load_workbook(cfg["chemins"]["classeur_indicateurs"], data_only=True)["Indicator Definitions"]
    defs = [wbd.cell(r, 3).value for r in range(4, wbd.max_row + 1) if wbd.cell(r, 3).value]
    doublons = pd.Series([cle(x) for x in defs]).duplicated().sum()
    res("IND-08", "Conformité", "Cadre global (feuille Definitions)", "Intitulés", "Aucun doublon",
        "0", f"{doublons} doublon(s) ; {len(defs)} lignes", "ÉCHEC" if doublons else "OK", "Mineure",
        "La ligne 28 duplique les visites de supervision au lieu de l'indicateur IDSR.")
    src_csv = csv.dropna(subset=["Indicator"])["Primary Data Source / Means of Verification"].astype(str)
    src_xls = [str(wbd.cell(r, 7).value) for r in range(4, wbd.max_row + 1) if wbd.cell(r, 3).value]
    n_odk_csv = sum("ODK" in x for x in src_csv)
    n_odk_xls = sum("ODK" in x for x in src_xls)
    res("IND-08", "Conformité", "Cadre global CSV vs Excel", "Version du cadre", "Une seule version de référence",
        "versions identiques", f"ODK cité {n_odk_csv} fois (CSV) contre {n_odk_xls} fois (Excel)",
        "ÉCHEC" if n_odk_csv != n_odk_xls else "OK", "Mineure", "Archiver le CSV ; la matrice maîtresse fait foi.")


# ==========================================================================
# E. Matrice de Dr Barclay
# ==========================================================================

def controles_barclay(cfg):
    wb = openpyxl.load_workbook(cfg["chemins"]["matrice_barclay"], data_only=True)
    ws = wb["Indicators"]
    rows = []
    for r in range(5, ws.max_row + 1):
        ident, lib, cible = ws.cell(r, 3).value, ws.cell(r, 4).value, ws.cell(r, 10).value
        freq = ws.cell(r, 9).value
        if ident is None or lib is None:
            continue
        rows.append(dict(ligne=r, id=ident, libelle=str(lib).strip(), cible=cible, freq=freq,
                         pilier=ws.cell(r, 2).value, actual=ws.cell(r, 11).value))
    df = pd.DataFrame(rows)
    dup = df[df.id.astype(str).duplicated(keep=False)]
    res("BAR-01", "Format", "Matrice de Dr Barclay", "Identifiants", "Identifiants uniques", "0 doublon",
        f"{dup.id.astype(str).nunique()} identifiants en double : {', '.join(sorted(dup.id.astype(str).unique()))}",
        "ÉCHEC" if len(dup) else "OK", "Majeure",
        "Les numéros sont stockés comme des nombres (1.10 devient 1.1) : utiliser des identifiants texte (SURV-10).")
    types = df.cible.map(lambda v: "texte" if isinstance(v, str) else ("1 (=100 % ?)" if v == 1 else "nombre"))
    n_txt, n_un = (types == "texte").sum(), (types == "1 (=100 % ?)").sum()
    res("BAR-02", "Format", "Matrice de Dr Barclay", "Cibles", "Cibles numériques homogènes", "un seul format",
        f"{n_txt} cibles en texte (« ≥90% », « Monitor »), {n_un} cibles codées 1", "ÉCHEC" if n_txt + n_un else "OK",
        "Majeure", "Séparer valeur cible (nombre) et unité ; écrire 100 % pour les pourcentages.")
    lib_dup = df[df.libelle.map(cle).duplicated(keep=False)]
    res("BAR-03", "Conformité", "Matrice de Dr Barclay", "Indicateurs dupliqués", "Aucun doublon d'intitulé", "0",
        "; ".join(f"{r.id} {r.libelle[:60]}" for r in lib_dup.itertuples()) or "aucun",
        "ÉCHEC" if len(lib_dup) else "OK", "Mineure", "Supprimer le doublon (LOG-14 = MEL-14).")
    s = wb["Summary "] if "Summary " in wb.sheetnames else None
    if s is not None:
        decl = {str(s.cell(r, 3).value).split(". ")[-1].strip(): s.cell(r, 4).value for r in range(7, 13)}
        reel = df.groupby("pilier").size()
        ecarts = []
        for k, v in reel.items():
            nom = str(k).split(". ")[-1].strip()
            d = next((dv for dk, dv in decl.items() if dk[:4].lower() == nom[:4].lower()), None)
            if d is not None and d != v:
                ecarts.append(f"{nom} {d} déclarés / {v} réels")
        res("BAR-04", "Cohérence interne", "Matrice de Dr Barclay", "Feuille Summary",
            "Effectifs du résumé = matrice", f"{len(df)} indicateurs", "; ".join(ecarts) or "conforme",
            "ÉCHEC" if ecarts else "OK", "Mineure", "Calculer le résumé par formule (NB.SI).")
    vide = df.actual.isna().all()
    res("BAR-05", "Exhaustivité", "Matrice de Dr Barclay", "Réalisations", "Colonnes Actual / % / Status renseignées",
        "renseignées", "vides" if vide else "partiellement", "INFO" if vide else "OK", "",
        "Alimenter automatiquement depuis la base (feuille saisie_hebdo).")
    df["rang"] = df.groupby("pilier").cumcount() + 1
    df["id_corrige"] = df.pilier.astype(str).str.extract(r"^(\d+)")[0] + "." + df.rang.astype(str)
    pm = df[df.libelle.str.contains("risk mapping", case=False)]
    for r in pm.itertuples():
        if str(r.freq).lower().startswith("annual") and isinstance(r.cible, (int, float)) and r.cible > 1:
            res("BAR-06", "Conformité", "Matrice de Dr Barclay", f"{r.id_corrige} (saisi {r.id}) {r.libelle}", "Fréquence compatible avec la cible",
                "mensuelle", f"annuelle pour une cible de {r.cible}", "ÉCHEC", "Mineure", "Passer en fréquence mensuelle.")


# ==========================================================================
# F. Plans de travail (budgets et cibles)
# ==========================================================================

ARGENT = re.compile(r"(?P<pre>.{0,3})\$\s?(?P<v>\d[\d,]*(?:\.\d+)?)(?P<post>.{0,12})", re.S)


def montant_detail(texte: str) -> float:
    """Somme des montants de lignes d'un texte de détail budgétaire (les coûts unitaires sont ignorés)."""
    total = 0.0
    jetons = []
    for m in re.finditer(r"\$\s?(\d[\d,]*(?:\.\d+)?)", texte):
        jetons.append((m.start(), m.end(), float(m.group(1).replace(",", ""))))
    for i, (d, fn, v) in enumerate(jetons):
        avant = texte[max(0, d - 4):d]
        apres = texte[fn:fn + 14]
        multiplicande = re.search(r"[X×x]\s*$", avant) is not None
        unitaire = re.match(r"\s*(/|per\b|=)", apres) is not None
        suivi_montant = re.match(r"\s{0,2}\$", apres) is not None
        if unitaire or suivi_montant or (multiplicande and re.match(r"\s*=", apres)):
            continue
        total += v
    # montants sans « $ » : après « = » (résultat) ou après « : » s'ils terminent l'élément
    for m in re.finditer(r"[=:]\s*(\d{1,3}(?:,\d{3})+)(?=\s*(?:$|[;.]|[A-Z]))", texte):
        total += float(m.group(1).replace(",", ""))
    return total


def couts_unitaires_suspects(texte: str, montant_ligne: float) -> list[str]:
    """Quantité x coût unitaire supérieure au montant de l'activité : coût unitaire mal saisi."""
    sortie = []
    for m in re.finditer(r"(\d[\d,]*)\s+[A-Za-z /-]{2,40}?\s[X×]\s?\$\s?(\d[\d,]*(?:\.\d+)?)(?!\s*(?:=|/|per))", texte):
        q, u = float(m.group(1).replace(",", "")), float(m.group(2).replace(",", ""))
        if montant_ligne and q * u > montant_ligne:
            sortie.append(m.group(0).strip())
    return sortie


def controles_plans(cfg, base):
    d = docx.Document(cfg["chemins"]["plan_dos"])
    lignes_act, totaux = {}, []
    for t in d.tables:
        for row in t.rows:
            cells = []
            for c in row.cells:
                tx = " ".join(c.text.split())
                if not cells or cells[-1] != tx:
                    cells.append(tx)
            if len(cells) >= 3 and re.match(r"^\d+\.\d+\.\d+$", cells[0]):
                cle_act = ".".join(cells[0].split(".")[:2])
                montant = nombre_fr(cells[-1].replace("$", "").replace(",", ""))
                detail = cells[-2] if len(cells) >= 4 else ""
                lignes_act.setdefault(cle_act, [])
                if (montant, detail) not in lignes_act[cle_act]:
                    lignes_act[cle_act].append((montant, detail))
            if len(cells) >= 2 and cells[-2].strip() == "Total":
                totaux.append(nombre_fr(cells[-1].replace("$", "").replace(",", "")))
    somme_pilier = {}
    for k, v in lignes_act.items():
        p = k.split(".")[0]
        somme_pilier[p] = somme_pilier.get(p, 0) + sum(m or 0 for m, _ in v)
        for m_, det in v:
            calc = montant_detail(det)
            if m_ and calc and abs(calc - m_) > 1:
                res("PT-03", "Cohérence interne", "Plan de travail DOS", f"Activité {k}",
                    "Détail des coûts = montant de l'activité", f"{m_:,.0f} $".replace(",", " "),
                    f"{calc:,.0f} $ (somme du détail)".replace(",", " "), "ÉCHEC", "Mineure",
                    "Corriger le détail ou le montant de l'activité.")
            elif m_ and calc:
                res("PT-03", "Cohérence interne", "Plan de travail DOS", f"Activité {k}",
                    "Détail des coûts = montant de l'activité", f"{m_:,.0f} $".replace(",", " "),
                    f"{calc:,.0f} $".replace(",", " "))
            for sus in couts_unitaires_suspects(det, m_):
                res("PT-07", "Plausibilité", "Plan de travail DOS", f"Activité {k}",
                    "Quantité x coût unitaire ≤ montant de l'activité", f"≤ {m_:,.0f} $".replace(",", " "),
                    f"« {sus} »", "ÉCHEC", "Mineure", "Coût unitaire mal saisi : corriger le libellé du détail.")
    for i, (p, s) in enumerate(sorted(somme_pilier.items(), key=lambda kv: int(kv[0]))):
        if i < len(totaux):
            ok = abs(s - totaux[i]) < 1
            res("PT-01", "Cohérence interne", "Plan de travail DOS", f"Pilier {p} : total",
                "Somme des activités = total du pilier", f"{totaux[i]:,.0f} $".replace(",", " "),
                f"{s:,.0f} $".replace(",", " "), "OK" if ok else "ÉCHEC", "Majeure",
                "Corriger le total du pilier.")
    s_lignes, s_totaux = sum(somme_pilier.values()), sum(t for t in totaux if t)
    res("PT-02", "Cohérence interne", "Plan de travail DOS", "Budget total",
        "Somme des piliers = 4 000 000 $ (budget par catégorie)", "4 000 000 $",
        f"lignes : {s_lignes:,.0f} $ ; totaux de piliers : {s_totaux:,.0f} $".replace(",", " "),
        "OK" if abs(s_lignes - 4_000_000) < 1 and abs(s_totaux - 4_000_000) < 1 else "ÉCHEC", "Majeure",
        "Réconcilier les tableaux d'activités et le budget par catégorie avant toute analyse coût-extrant.")
    texte = "\n".join(p.text for p in d.paragraphs) + "\n" + "\n".join(
        c.text for t in d.tables for row in t.rows for c in row.cells)
    hebdo = sorted({int(x) for x in re.findall(r"(\d+)\s+weekly (?:coordination )?updates", texte)})
    res("PT-04", "Cohérence interne", "Plan de travail DOS", "Mises à jour hebdomadaires", "Une seule cible",
        "1 valeur", " et ".join(map(str, hebdo)), "ÉCHEC" if len(hebdo) > 1 else "OK", "Mineure",
        "Arrêter une cible unique (COORD-09).")
    # Tranche 5 : budget
    d5 = docx.Document(cfg["chemins"]["plan_tranche5"])
    piliers5, total5 = [], None
    for t in d5.tables:
        for row in t.rows:
            cells = [" ".join(c.text.split()) for c in row.cells]
            if len(cells) >= 2 and cells[0].startswith("Pillar") and "$" in cells[1]:
                piliers5.append(nombre_fr(cells[1].replace("$", "").replace(",", "")))
            if cells and cells[0].upper() == "TOTAL" and len(cells) > 1:
                total5 = nombre_fr(cells[1].replace("$", "").replace(",", ""))
    if piliers5 and total5:
        ok = abs(sum(piliers5) - total5) < 1
        res("PT-06", "Cohérence interne", "Plan Tranche 5", "Budget", "Somme des piliers = total",
            f"{total5:,.0f} $".replace(",", " "), f"{sum(piliers5):,.0f} $".replace(",", " "),
            "OK" if ok else "ÉCHEC", "Majeure")
    # PT-05 cibles DOS vs Tranche 5
    ref = lire_table(base, "ref_indicateurs")
    def cible_num(x):
        x = str(x).strip()
        return nombre_fr(x) if re.fullmatch(r"[\d\s\u202f\u00a0]+", x) else None
    ecarts = []
    for r in ref.itertuples():
        a, b = cible_num(r.cible_dos), cible_num(r.cible_tranche5)
        if a and b and r.type == "Nombre":
            ecarts.append((r.id_indicateur, a, b, b / a))
    if ecarts:
        txt = "; ".join(f"{i} ×{x:.1f}".replace(".", ",") for i, _, _, x in ecarts)
        res("PT-05", "Cohérence externe", "Plan DOS vs Tranche 5", f"{len(ecarts)} cibles comparables",
            "Un seul jeu de cibles de référence par période", "cibles identiques", txt, "ÉCHEC", "Critique",
            "Décider quelle cible s'applique à quelle période et à quel financement avant le prochain rapport DOS.")


# ==========================================================================
# G. Fiche SBC et observations manuelles
# ==========================================================================

def controles_sbc(cfg):
    wb = openpyxl.load_workbook(cfg["chemins"]["fiche_sbc"])
    ws = wb[wb.sheetnames[0]]
    entetes = " ".join(str(ws.cell(10, c).value or "") for c in range(1, ws.max_column + 1)).lower()
    attendus = {
        "vérification des alertes en 24 h (SURV-06)": "vérifi",
        "rumeurs enregistrées et traitées en 48 h (CREC-05/06)": "rumeur",
        "contacts ayant terminé 21 jours (SURV-21)": "21",
        "contacts symptomatiques référés (SURV-26)": "référ",
        "décès communautaires notifiés (SURV-30)": "notifi",
    }
    manquants = [k for k, mot in attendus.items() if mot not in entetes]
    res("SBC-01", "Conformité", "Fiche SBC journalière", "Champs requis par les indicateurs",
        "Chaque indicateur alimenté par la fiche a son champ", f"{len(attendus)} champs",
        "absents : " + "; ".join(manquants) if manquants else "tous présents", "ÉCHEC" if manquants else "OK", "Majeure",
        "Ajouter les champs « NOUVEAU » du canevas (feuille sbc_synthese_as).")
    fen = []
    if "investiguées dans 24h" in entetes or "investiguées dans 24" in entetes:
        fen.append("investigation en 24 h dans la fiche contre 7 jours dans le cadre (SURV-08)")
    if "levées le même jour" in entetes:
        fen.append("réticences levées le jour même contre rumeurs traitées en 48 h (CREC-06)")
    res("SBC-02", "Conformité", "Fiche SBC journalière", "Fenêtres temporelles", "Même délai que la définition",
        "délais identiques", "; ".join(fen) or "conformes", "ÉCHEC" if fen else "OK", "Majeure",
        "Harmoniser sur 24 h (investigation) et 48 h (rumeurs, réticences).")
    multi = sum(1 for r in range(12, 34) if str(ws.cell(r, 5).value or "").count(":") >= 2)
    res("SBC-03", "Format", "Fiche SBC journalière", "Colonne « technique de communication »",
        "Une valeur numérique par cellule", "1 valeur", f"{multi} lignes à plusieurs comptages en texte libre",
        "ÉCHEC" if multi else "OK", "Mineure", "Une colonne numérique par technique (canevas sbc_activites).")
    obs = pd.read_csv(cfg["chemins"]["referentiels"] / "observations_manuelles.csv")
    for r in obs.itertuples():
        res(r.id_regle, "Observation manuelle", r.source, r.objet, r.description,
            "" if pd.isna(r.valeur_reference) else str(r.valeur_reference), str(r.valeur_observee), r.statut,
            "" if pd.isna(r.gravite) else r.gravite, "" if pd.isna(r.recommandation) else r.recommandation)



# ==========================================================================
# H. Weekly Updates CGPP et série nationale
# ==========================================================================

def controles_weekly(cfg, base):
    inter = cfg["chemins"]["intermediaire"]
    if not (inter / "weekly_national.csv").exists():
        return
    lire = lambda nom: pd.read_csv(inter / nom) if (inter / nom).exists() else pd.DataFrame()  # noqa: E731
    nat, ops, puces, resumes = (lire("weekly_national.csv"), lire("weekly_operations.csv"),
                                lire("weekly_puces.csv"), lire("weekly_resumes.csv"))
    tol = cfg["seuils"]["tolerance_pourcentage"] + 0.05      # arrondi de publication à une décimale
    nom = lambda d: f"Weekly Update {pd.to_datetime(d):%d/%m}"  # noqa: E731

    for d, g in nat.groupby("date_weekly"):
        src, no = nom(d), (int(g.sitrep_no.dropna().iloc[0]) if g.sitrep_no.notna().any() else "?")
        prov, tot = g[g.province != "Total"], g[g.province == "Total"]
        for col, lib in [("cas_cumul", "cas confirmés"), ("deces_cumul", "décès confirmés"),
                         ("nouveaux_cas_24h", "nouveaux cas (24 h)"), ("zs_touchees", "ZS touchées")]:
            somme, total = prov[col].sum(), (tot[col].iloc[0] if len(tot) else None)
            ok = egal(somme, total)
            res("WU-01", "Cohérence interne", src, f"Tableau national (SitRep {no}) : {lib}",
                "La somme des provinces égale la ligne Total", f(somme) + " (somme des provinces)", f(total) + " (publié)",
                "OK" if ok else "ÉCHEC", "Mineure", "Corriger la ligne Total ou la province en cause avant l'envoi.")
        ecarts = []
        for r in g.itertuples():
            if r.cas_cumul and pd.notna(r.letalite_publiee):
                calc = round(100 * r.deces_cumul / r.cas_cumul, 1)
                if abs(calc - r.letalite_publiee) > tol:
                    ecarts.append(f"{r.province} : {f(r.letalite_publiee)} % publié, {f(calc)} % recalculé")
        res("WU-02", "Cohérence interne", src, f"Tableau national (SitRep {no}) : létalité",
            "Létalité = décès / cas, par province", "écart ≤ 0,15 point", "; ".join(ecarts) or "conformes",
            "ÉCHEC" if ecarts else "OK", "Mineure", "Recalculer la létalité à partir des cumuls.")
        r0 = resumes[resumes.date_weekly == d]
        if len(r0) and len(tot):
            r0, t0 = r0.iloc[0], tot.iloc[0]
            diff = [x for x, a, b in [("cas", r0.cas, t0.cas_cumul), ("décès", r0.deces, t0.deces_cumul)]
                    if pd.notna(a) and not egal(a, b)]
            res("WU-03", "Cohérence interne", src, "Résumé exécutif contre tableau national",
                "Les cumuls du résumé sont ceux du tableau", f"{f(t0.cas_cumul)} cas ; {f(t0.deces_cumul)} décès",
                f"{f(r0.cas)} cas ; {f(r0.deces)} décès", "ÉCHEC" if diff else "OK", "Mineure",
                "Reprendre les chiffres du tableau dans le résumé.")

    if len(ops):
        for d, g in ops.groupby("date_weekly"):
            src = nom(d)
            for lig, h in g.groupby("ligne", sort=False):
                ips, tot = h[h.ip != "Total"], h[h.ip == "Total"]
                if not len(tot) or "rate" in lig.lower():
                    continue
                somme, total = ips.valeur.sum(), tot.valeur.iloc[0]
                ok = egal(somme, total)
                res("WU-04", "Cohérence interne", src, f"Tableau opérationnel : {lig}",
                    "La colonne Total égale la somme des IP", f(somme) + " (somme des IP)", f(total) + " (publié)",
                    "OK" if ok else "ÉCHEC", "Majeure",
                    "Calculer la colonne Total par formule (SOMME) à partir des colonnes des IP.")
            cle_col = lambda x: x.set_index(["ip", "province"]).valeur  # noqa: E731
            att = cle_col(g[g.cle_ligne == "expectedrecos"])
            act = cle_col(g[g.cle_ligne == "recoswhoworked"])
            taux = cle_col(g[g.cle_ligne == "recoparticipationrate"])
            for k, v in taux.items():
                if k in att.index and att[k]:
                    calc = round(100 * act[k] / att[k], 1)
                    ok = abs(calc - v) <= tol
                    lib = "total" if k[0] == "Total" else f"{k[0]} {k[1]}"
                    res("WU-05", "Cohérence interne", src, f"Taux de participation des RECO ({lib})",
                        "Taux = RECO actifs / RECO attendus", f"{f(calc)} % ({f(act[k])}/{f(att[k])})", f"{f(v)} %",
                        "OK" if ok else "ÉCHEC", "Majeure" if k[0] == "Total" else "Mineure",
                        "Recalculer le taux à partir des totaux, jamais en moyennant les taux des IP.")
            attendus = set(cfg.get("partenaires", {}).keys())
            presents = set(g.ip) - {"Total"}
            manquants = sorted(attendus - presents)
            res("WU-06", "Exhaustivité", src, "Partenaires présents dans le tableau opérationnel",
                "Tous les IP du consortium", ", ".join(sorted(attendus)), ", ".join(sorted(presents)) +
                (f" ; absents : {', '.join(manquants)}" if manquants else ""), "ÉCHEC" if manquants else "OK",
                "Majeure", "Obtenir les données d'IMC et d'IRC ou indiquer explicitement « non rapporté ».")
        # fraîcheur des données opérationnelles au moment du SitRep
        meta = lire_table(base, "sitrep_meta")
        if len(meta):
            d_sitrep = pd.to_datetime(meta.date_rapport).max()
            fin = pd.to_datetime(ops.fin).max()
            retard = (d_sitrep - fin).days
            res("WU-09", "Ponctualité", nom(ops.date_weekly.max()), "Âge des données opérationnelles CGPP",
                "Données de 7 jours au plus à la date du SitRep", f"≤ 7 jours (SitRep du {d_sitrep:%d/%m})",
                f"{retard} jours (données jusqu'au {fin:%d/%m})", "ÉCHEC" if retard > 7 else "OK", "Majeure",
                "Transmettre le tableau consolidé chaque semaine, au plus tard le lundi.")

    if len(puces):
        dates = sorted(puces.date_weekly.unique())
        for d0, d1 in zip(dates, dates[1:]):
            a = puces[(puces.date_weekly == d0) & puces.section.str.contains("Other Disease", case=False, na=False)]
            b = puces[(puces.date_weekly == d1) & puces.section.str.contains("Other Disease", case=False, na=False)]
            identique = len(a) > 2 and set(a.texte) == set(b.texte)
            res("WU-07", "Actualité", nom(d1), "Section « autres épidémies »",
                "Mise à jour à chaque rapport", f"différente du {pd.to_datetime(d0):%d/%m}",
                f"identique au rapport du {pd.to_datetime(d0):%d/%m} ({len(b)} paragraphes)" if identique else "mise à jour",
                "ÉCHEC" if identique else "OK", "Mineure",
                "Dater la semaine épidémiologique et mettre la section à jour, ou l'omettre si rien n'a changé.")
        for d, g in puces.groupby("date_weekly"):
            dbl = g[g.duplicated(["section", "texte"])]
            res("WU-08", "Format", nom(d), "Puces répétées dans une même section", "Aucune répétition", "0",
                f"{len(dbl)} : " + " ; ".join(t[:60] for t in dbl.texte) if len(dbl) else "0",
                "ÉCHEC" if len(dbl) else "OK", "Mineure", "Relire le rapport avant envoi.")

    # série nationale : cumuls croissants d'un SitRep à l'autre, provinces = total national
    hist = lire_table(base, "historique")
    h = hist[hist.source.astype(str).str.contains("SitRep", na=False) & hist.indicateur.isin(["cas_cumul", "deces_cumul"])]
    ep = lire_table(base, "epi_province")
    meta = lire_table(base, "sitrep_meta")
    lignes = [dict(date=pd.to_datetime(r.date_reference), territoire=r.territoire, indicateur=r.indicateur,
                   valeur=r.valeur) for r in h.itertuples()]
    dates_meta = dict(zip(meta.sitrep_no, pd.to_datetime(meta.date_rapport)))
    for r in ep.itertuples():
        for ind in ("cas_cumul", "deces_cumul"):
            lignes.append(dict(date=dates_meta.get(r.sitrep_no), territoire=r.province, indicateur=ind,
                               valeur=getattr(r, ind)))
    for r in meta.itertuples():
        lignes.append(dict(date=pd.to_datetime(r.date_rapport), territoire="National", indicateur="cas_cumul", valeur=r.cas_cumul))
        lignes.append(dict(date=pd.to_datetime(r.date_rapport), territoire="National", indicateur="deces_cumul", valeur=r.deces_cumul))
    serie = pd.DataFrame(lignes).dropna(subset=["date"]).drop_duplicates(["date", "territoire", "indicateur"], keep="last")
    baisses = []
    for (terr, ind), g in serie.sort_values("date").groupby(["territoire", "indicateur"]):
        v = g.valeur.astype(float).values
        for i in range(1, len(v)):
            if v[i] < v[i - 1]:
                baisses.append(f"{terr} {ind.replace('_cumul', '')} : {f(v[i - 1])} puis {f(v[i])} "
                               f"({g.date.iloc[i]:%d/%m})")
    res("TS-05", "Cohérence temporelle", "Série nationale (SitReps 085 à 149)", "Cumuls par province et national",
        "Un cumul ne baisse jamais d'un SitRep à l'autre", f"{serie.date.nunique()} dates", "; ".join(baisses) or
        "aucune baisse", "ÉCHEC" if baisses else "OK", "Majeure", "Signaler tout reclassement en note du SitRep.")
    ecarts = []
    for (dt, ind), g in serie.groupby(["date", "indicateur"]):
        tot = g[g.territoire == "National"].valeur
        pr = g[g.territoire != "National"].valeur.astype(float).sum()
        if len(tot) and len(g) > 2 and not egal(pr, tot.iloc[0]):
            ecarts.append(f"{dt:%d/%m} {ind.replace('_cumul', '')} : provinces {f(pr)}, national {f(tot.iloc[0])}")
    res("TS-06", "Cohérence interne", "Série nationale (SitReps 085 à 149)", "Somme des provinces et total national",
        "Égalité à chaque date", "écart nul", "; ".join(ecarts) or "conformes", "ÉCHEC" if ecarts else "OK",
        "Mineure", "Vérifier la source (le SitRep 085 n'est connu que par le plan DOS).")

# ==========================================================================
# Rapport
# ==========================================================================

def famille(source: str) -> str:
    s = str(source)
    if s.startswith("SitRep"):
        return "SitReps nationaux"
    if s.startswith(("CR", "Comptes rendus")):
        return "Comptes rendus de coordination"
    if s.startswith(("Classeur", "Cadre global")):
        return "Classeur hebdomadaire et cadre global"
    if s.startswith("Matrice"):
        return "Matrice de Dr Barclay"
    if s.startswith(("Plan", "Plan DOS")):
        return "Plans de travail (DOS, Tranche 5)"
    if s.startswith("Fiche SBC"):
        return "Fiche SBC journalière"
    if s.startswith("Weekly Update"):
        return "Weekly Updates CGPP"
    if s.startswith("Série nationale"):
        return "SitReps nationaux"
    return "Autre"


def main(cfg: dict | None = None):
    cfg = cfg or charger_config()
    base = cfg["chemins"]["base"]
    R.clear()
    for etape in (controles_sitreps, controles_reunions, controles_classeur):
        etape(cfg, base)
    controles_barclay(cfg)
    controles_plans(cfg, base)
    controles_weekly(cfg, base)
    controles_sbc(cfg)

    df = pd.DataFrame(R)
    df.insert(0, "famille", df.source.map(famille))
    ordre_g = {"Critique": 0, "Majeure": 1, "Mineure": 2, "": 3}
    df["_o"] = df.gravite.map(ordre_g).fillna(3)
    anomalies = df[df.statut.isin(["ÉCHEC", "À VÉRIFIER"])].sort_values(["_o", "id_regle"]).drop(columns="_o")
    df = df.drop(columns="_o")
    synthese = (df.assign(n=1).pivot_table(index="famille", columns="statut", values="n", aggfunc="sum", fill_value=0)
                .reset_index())
    par_gravite = anomalies.groupby("gravite").size().reindex(["Critique", "Majeure", "Mineure"], fill_value=0)

    jour = date.today().isoformat()
    sortie = cfg["chemins"]["sorties_qc"] / f"rapport_qc_{jour}.xlsx"
    pcts = pd.read_csv(cfg["chemins"]["intermediaire"] / "pourcentages_texte_sitreps.csv")
    with pd.ExcelWriter(sortie, engine="xlsxwriter") as xw:
        resume = pd.DataFrame([
            ("Contrôles exécutés", len(df)),
            ("Conformes", int((df.statut == "OK").sum())),
            ("Échecs", int((df.statut == "ÉCHEC").sum())),
            ("À vérifier", int((df.statut == "À VÉRIFIER").sum())),
            ("Informations", int((df.statut == "INFO").sum())),
            ("Anomalies critiques", int(par_gravite["Critique"])),
            ("Anomalies majeures", int(par_gravite["Majeure"])),
            ("Anomalies mineures", int(par_gravite["Mineure"])),
        ], columns=["Mesure", "Valeur"])
        resume.to_excel(xw, sheet_name="synthese", index=False, startrow=0)
        synthese.to_excel(xw, sheet_name="synthese", index=False, startrow=len(resume) + 3)
        anomalies.to_excel(xw, sheet_name="anomalies", index=False)
        df.to_excel(xw, sheet_name="tous_les_controles", index=False)
        pcts.to_excel(xw, sheet_name="pourcentages_sitreps", index=False)
        for nom in ["synthese", "anomalies", "tous_les_controles", "pourcentages_sitreps"]:
            ws = xw.sheets[nom]
            ws.freeze_panes(1, 0)
            ws.set_column(0, 0, 14)
            ws.set_column(1, 3, 28)
            ws.set_column(4, 8, 40)
    # suivi des anomalies dans la base
    suivi = lire_table(base, "qc_suivi")
    nouv = anomalies.assign(id_anomalie=anomalies.id_regle + " | " + anomalies.source + " | " + anomalies.objet,
                            date_detection=pd.Timestamp(jour), regle=anomalies.id_regle, statut="Ouverte",
                            responsable=None, date_resolution=None)
    nouv = nouv.reindex(columns=[c[0] for c in SCHEMA["qc_suivi"]["colonnes"]])
    if not suivi.empty:
        nouv = nouv[~nouv.id_anomalie.isin(set(suivi.id_anomalie))]
    morceaux = [x for x in (suivi, nouv) if not x.empty]
    tout = pd.concat(morceaux, ignore_index=True) if len(morceaux) > 1 else (morceaux[0] if morceaux else nouv)
    ecrire_tables(base, {"qc_suivi": tout}, cfg["chemins"].get("canevas_base"))
    LOG.info("%d contrôles : %d conformes, %d échecs, %d à vérifier (critiques %d, majeures %d, mineures %d)",
             len(df), (df.statut == "OK").sum(), (df.statut == "ÉCHEC").sum(), (df.statut == "À VÉRIFIER").sum(),
             par_gravite["Critique"], par_gravite["Majeure"], par_gravite["Mineure"])
    LOG.info("Rapport : %s", sortie)
    return df, anomalies, sortie


if __name__ == "__main__":
    main()
