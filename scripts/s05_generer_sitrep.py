"""
s05_generer_sitrep.py - Génération automatique du SitRep CGPP (Word puis PDF) à partir de la base Excel.

Entrées : base Excel (s01, s02, s02b), résultats des analyses (s04), dernier rapport de contrôle qualité (s03),
          canevas Word modeles/canevas_sitrep_CGPP.docx (balises docxtpl).
Sorties : sorties/sitrep/SitRep_CGPP_AAAA-MM-JJ.docx et .pdf

Usage   : python scripts/s05_generer_sitrep.py              (dernier SitRep national de la base)
          python scripts/s05_generer_sitrep.py 2026-10-10   (date des données)

Contenu : situation nationale (série des SitReps, y compris ceux cités par les Weekly Updates), foyers et ZS du
dispositif, données opérationnelles consolidées des IP (dernier Weekly Update, totaux et taux recalculés),
activités, performance contre les cibles, sécurité, défis et actions (feuille sitrep_notes), autres épidémies,
points d'action des réunions et qualité des données. Les textes sont rédigés par des règles à partir des chiffres :
ils se relisent avant diffusion.
"""
from __future__ import annotations

import json
import sys
from datetime import datetime
from pathlib import Path

import docx
import pandas as pd
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor
from docx.text.paragraph import Paragraph
from docxtpl import DocxTemplate, InlineImage

sys.path.insert(0, str(Path(__file__).resolve().parent))
from commun import charger_config, journal, lire_table, vers_pdf  # noqa: E402

LOG = journal("sitrep")
VERT, JAUNE, ROUGE, GRIS, BLEU_PALE, BLANC = "D9F2D9", "FFF2CC", "F8D7D7", "EDEDED", "DCE6F2", "FFFFFF"
BLEU = "1F3864"
MOIS = ["janvier", "février", "mars", "avril", "mai", "juin", "juillet", "août", "septembre", "octobre",
        "novembre", "décembre"]
PLACE_OPERATIONS = "[[TABLEAU_OPERATIONS]]"


# --------------------------------------------------------------------------
# Mise en forme
# --------------------------------------------------------------------------

def n(x, dec=0, signe=False):
    """Nombre au format français (espace fine insécable pour les milliers)."""
    if x is None or (isinstance(x, float) and pd.isna(x)):
        return "ND"
    s = f"{x:+,.{dec}f}" if signe else f"{x:,.{dec}f}"
    return s.replace(",", " ").replace(".", ",")


def pct(x, dec=1):
    return "ND" if x is None or pd.isna(x) else n(x, dec) + " %"


def jour_mois(d, annee=False):
    d = pd.to_datetime(d)
    return f"{d.day}{'er' if d.day == 1 else ''} {MOIS[d.month - 1]}" + (f" {d.year}" if annee else "")


def couleur_taux(v, vert=95, jaune=90):
    if v is None or pd.isna(v):
        return GRIS
    return VERT if v >= vert else (JAUNE if v >= jaune else ROUGE)


def taux(num, den):
    return 100 * num / den if den and pd.notna(den) and pd.notna(num) else None


def liste_fr(elements):
    el = [str(e) for e in elements if str(e)]
    return el[0] if len(el) == 1 else (", ".join(el[:-1]) + " et " + el[-1] if el else "")


LIBELLES_OPS = {"People sensitized (CREC)": "personnes sensibilisées", "Expected RECOs": "RECO attendus",
                "RECOs who worked": "RECO actifs", "Number of reports received": "rapports reçus",
                "Sick alerts detected": "alertes détectées", "Sick alerts reported": "alertes notifiées",
                "Contacts under followup (cumulative)": "contacts suivis", "Contacts actually seen": "contacts vus",
                "Resistances identified": "résistances identifiées", "Resistances resolved": "résistances levées",
                "Community deaths swabbed": "décès prélevés", "Safe and Dignified Burials": "EDS"}


def ecarts_publies(cfg, fin) -> list[str]:
    """Totaux et taux publiés dans le tableau du Weekly Update qui diffèrent du recalcul."""
    f = cfg["chemins"]["intermediaire"] / "weekly_operations.csv"
    if not f.exists():
        return []
    o = pd.read_csv(f)
    o = o[pd.to_datetime(o.fin) == fin]
    sortie = []
    for lig, g in o.groupby("ligne", sort=False):
        ips, tot = g[g.ip != "Total"], g[g.ip == "Total"]
        if not len(tot) or pd.isna(tot.valeur.iloc[0]) or "rate" in lig.lower():
            continue
        pub, somme = float(tot.valeur.iloc[0]), float(ips.valeur.sum())
        if abs(pub - somme) > 0.5:
            sans = [r.ip for r in ips.itertuples() if abs(somme - r.valeur - pub) < 0.5]
            sortie.append(f"{n(pub)} {LIBELLES_OPS.get(lig, lig)}" + (f" (sans la colonne {sans[0]})" if sans else "") +
                          f" au lieu de {n(somme)}")
    att = o[(o.cle_ligne == "expectedrecos") & (o.ip == "Total")].valeur
    act = o[(o.cle_ligne == "recoswhoworked") & (o.ip == "Total")].valeur
    tx = o[(o.cle_ligne == "recoparticipationrate") & (o.ip == "Total")].valeur
    if len(att) and len(act) and len(tx) and att.iloc[0]:
        calc = 100 * act.iloc[0] / att.iloc[0]
        if abs(calc - tx.iloc[0]) > 0.2:
            sortie.append(f"un taux de participation des RECO de {pct(tx.iloc[0])} au lieu de {pct(calc)}")
    return sortie


# --------------------------------------------------------------------------
# Tableau des données opérationnelles (colonnes variables : inséré après le rendu docxtpl)
# --------------------------------------------------------------------------

def _xml(tag, **attrs):
    el = OxmlElement(tag)
    for k, v in attrs.items():
        el.set(qn(f"w:{k}"), str(v))
    return el


ORDRE_TBLPR = ["tblStyle", "tblpPr", "tblOverlap", "bidiVisual", "tblStyleRowBandSize", "tblStyleColBandSize", "tblW",
               "jc", "tblCellSpacing", "tblInd", "tblBorders", "shd", "tblLayout", "tblCellMar", "tblLook"]
ORDRE_TRPR = ["cnfStyle", "divId", "gridBefore", "gridAfter", "wBefore", "wAfter", "cantSplit", "trHeight", "tblHeader",
              "tblCellSpacing", "jc", "hidden"]
ORDRE_TCPR = ["cnfStyle", "tcW", "gridSpan", "hMerge", "vMerge", "tcBorders", "shd", "noWrap", "tcMar", "textDirection",
              "tcFitText", "vAlign", "hideMark", "headers", "cellIns", "cellDel", "cellMerge", "tcPrChange"]


def ordonner_document(doc):
    """Ordre des propriétés de tableau conforme au schéma dans tout le document (la balise cellbg de docxtpl
    insère la trame en tête de tcPr, ce que Word peut signaler comme contenu illisible)."""
    corps = doc.element.body
    for tag, ordre in (("w:tblPr", ORDRE_TBLPR), ("w:trPr", ORDRE_TRPR), ("w:tcPr", ORDRE_TCPR)):
        for el in corps.iter(qn(tag)):
            _ordonner(el, ordre)


def _ordonner(parent, ordre):
    """Remet les enfants d'un élément de propriétés dans l'ordre imposé par le schéma OOXML (Word y est strict)."""
    rang = {qn(f"w:{t}"): i for i, t in enumerate(ordre)}
    enfants = list(parent)
    for e in enfants:
        parent.remove(e)
    for e in sorted(enfants, key=lambda e: rang.get(e.tag, len(ordre))):
        parent.append(e)


def _fond(cellule, couleur):
    tcpr = cellule._tc.get_or_add_tcPr()
    tcpr.append(_xml("w:shd", val="clear", color="auto", fill=couleur))


def _ecrire(cellule, texte, gras=False, couleur="0B0B0B", gauche=False, taille=8.5):
    p = cellule.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT if gauche else WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(str(texte))
    r.font.name, r.font.size, r.bold = "Arial", Pt(taille), gras
    r.font.color.rgb = RGBColor.from_string(couleur)
    r._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:cs"), "Arial")


def inserer_tableau(doc, marque, entetes, lignes, largeurs_mm):
    """Remplace le paragraphe `marque` par un tableau (en-tête bleu, couleurs par cellule)."""
    cible = next((p for p in doc.paragraphs if p.text.strip() == marque), None)
    if cible is None:
        return False
    t = doc.add_table(rows=1 + len(lignes), cols=len(entetes))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    tblpr = t._tbl.tblPr
    total = int(sum(largeurs_mm) * 56.693)
    for el in tblpr.findall(qn("w:tblW")):
        tblpr.remove(el)
    tblpr.append(_xml("w:tblW", w=total, type="dxa"))     # tblLayout « fixed » : déjà posé par autofit = False
    bords = _xml("w:tblBorders")
    for cote in ("top", "left", "bottom", "right", "insideH", "insideV"):
        bords.append(_xml(f"w:{cote}", val="single", sz=4, space=0, color="BFBFBF"))
    tblpr.append(bords)
    marges = _xml("w:tblCellMar")
    for cote, v in (("top", 30), ("left", 80), ("bottom", 30), ("right", 80)):
        marges.append(_xml(f"w:{cote}", w=v, type="dxa"))
    tblpr.append(marges)
    for gc, w in zip(t._tbl.tblGrid.findall(qn("w:gridCol")), largeurs_mm):
        gc.set(qn("w:w"), str(int(w * 56.693)))
    for ligne in t.rows:
        for c, w in zip(ligne.cells, largeurs_mm):
            c.width = Mm(w)
    for j, e in enumerate(entetes):
        c = t.rows[0].cells[j]
        _fond(c, BLEU)
        _ecrire(c, e, gras=True, couleur="FFFFFF", gauche=(j == 0), taille=8)
    t.rows[0]._tr.get_or_add_trPr().append(_xml("w:tblHeader", val="true"))
    for i, ligne in enumerate(lignes, start=1):
        for j, v in enumerate(ligne["cellules"]):
            c = t.rows[i].cells[j]
            _ecrire(c, v, gras=(j == len(entetes) - 1) or ligne.get("gras", False), gauche=(j == 0))
            coul = ligne.get("couleurs", {}).get(j) or (BLEU_PALE if j == len(entetes) - 1 else None)
            if coul:
                _fond(c, coul)
    # garder le tableau d'un seul tenant (chaque ligne liée à la suivante, aucune ligne coupée)
    for ligne in t.rows:
        ligne._tr.get_or_add_trPr().append(_xml("w:cantSplit", val="true"))
    for ligne in t.rows[:-1]:
        for c in ligne.cells:
            for par in c.paragraphs:
                par.paragraph_format.keep_with_next = True
    precedent = cible._p.getprevious()
    if precedent is not None and precedent.tag == qn("w:p"):
        Paragraph(precedent, cible._parent).paragraph_format.keep_with_next = True
    _ordonner(tblpr, ORDRE_TBLPR)
    for ligne in t.rows:
        _ordonner(ligne._tr.get_or_add_trPr(), ORDRE_TRPR)
        for c in ligne.cells:
            _ordonner(c._tc.get_or_add_tcPr(), ORDRE_TCPR)
    cible._p.addnext(t._tbl)
    cible._p.getparent().remove(cible._p)
    return True


# --------------------------------------------------------------------------
# Contexte du SitRep
# --------------------------------------------------------------------------

def contexte(cfg, date_cible=None):
    base = cfg["chemins"]["base"]
    meta = lire_table(base, "sitrep_meta").sort_values("sitrep_no")
    meta["date_rapport"] = pd.to_datetime(meta["date_rapport"])
    if date_cible:
        meta = meta[meta.date_rapport <= pd.to_datetime(date_cible)]
    m1 = meta.iloc[-1]
    m0 = meta.iloc[-2] if len(meta) > 1 else m1
    no1, no0 = int(m1.sitrep_no), int(m0.sitrep_no)
    d1 = m1.date_rapport
    res = json.loads((cfg["chemins"]["sorties_analyses"] / "resultats.json").read_text(encoding="utf-8"))
    nat, cg, cv, q = res["national"], res["cgpp"], res["couverture"], res.get("quinzaine", {})
    ivs = res.get("intervalles", [])
    jours = nat["jours"]
    d_ref = pd.to_datetime(q["date_ref"]) if q else None

    prov = lire_table(base, "epi_province")
    zs = lire_table(base, "epi_zs")
    contacts = lire_table(base, "contacts_province")
    p1 = prov[prov.sitrep_no == no1].set_index("province")
    p0 = prov[prov.sitrep_no == no0].set_index("province")
    c1 = contacts[contacts.sitrep_no == no1].set_index("province")
    zs1, zs0 = zs[zs.sitrep_no == no1], zs[zs.sitrep_no == no0]

    # ---------------- 1. provinces
    lignes_prov = []
    for p, r in p1.sort_values("cas_cumul", ascending=False).iterrows():
        h3 = r.cas_cumul - (p0.loc[p].cas_cumul if p in p0.index else 0)
        h14 = q.get("provinces", {}).get(p) if q else None
        tx, statut_c = None, "-"
        if p in c1.index:
            rc = c1.loc[p]
            statut_c = "ND"
            if rc.statut_donnees == "OK" and rc.contacts_suivis:
                tx = 100 * rc.contacts_vus / rc.contacts_suivis
                statut_c = pct(tx)
        lignes_prov.append(dict(nom=p, nouveaux=n(r.nouveaux_cas_24h), cas=n(r.cas_cumul),
                                h14=n(h14, signe=True) if h14 else "0", hausse=n(h3, signe=True) if h3 else "0",
                                deces=n(r.deces_cumul), letalite=pct(100 * r.deces_cumul / r.cas_cumul),
                                zs=f"{int(r.zs_touchees)}/{int(r.zs_total)}", contacts=statut_c,
                                couleur_contacts=couleur_taux(tx) if p in c1.index else BLANC))
    c_ok = c1[c1.statut_donnees == "OK"]
    nd_prov = list(c1[c1.statut_donnees == "ND"].index)
    taux_tot = taux(c_ok.contacts_vus.sum(), c_ok.contacts_suivis.sum()) if len(c_ok) else None
    tot = dict(nouveaux=n(p1.nouveaux_cas_24h.sum()), cas=n(p1.cas_cumul.sum()),
               h14=n(q.get("cas"), signe=True) if q else "-",
               hausse=n(p1.cas_cumul.sum() - p0.cas_cumul.sum(), signe=True), deces=n(p1.deces_cumul.sum()),
               letalite=pct(100 * p1.deces_cumul.sum() / p1.cas_cumul.sum()),
               zs=f"{int(p1.zs_touchees.sum())}/{int(p1.zs_total.sum())}",
               contacts=pct(taux_tot) + (" (hors ND)" if nd_prov else ""))

    # ---------------- 2. ZS CGPP
    zc = cfg["zones_cgpp"]
    lignes_zs, tz = [], dict(cas=0, hausse=0, nouveaux=0, deces=0)
    nouv_cgpp = post_cgpp = 0
    for nom, info in zc.items():
        a, b = zs1[zs1.zone_sante == nom], zs0[zs0.zone_sante == nom]
        cas = float(a.cas_cumul.iloc[0]) if len(a) and pd.notna(a.cas_cumul.iloc[0]) else 0
        dec = float(a.deces_cumul.iloc[0]) if len(a) else 0
        nv = float(a.nouveaux_cas_24h.fillna(0).iloc[0]) if len(a) else 0
        pm = float(a.deces_communautaires_24h.fillna(0).iloc[0]) if len(a) else 0
        prec = float(b.cas_cumul.iloc[0]) if len(b) and pd.notna(b.cas_cumul.iloc[0]) else 0
        h = cas - prec
        nouv_cgpp += nv
        post_cgpp += pm
        tz["cas"] += cas; tz["hausse"] += h; tz["nouveaux"] += nv; tz["deces"] += dec
        lignes_zs.append(dict(nom=nom, province=info["province"], ip=info["ip"], cas=n(cas), _cas=cas, _h=h,
                              hausse=n(h, signe=True) if h else "0", nouveaux=n(nv), deces=n(dec),
                              letalite=pct(100 * dec / cas) if cas else "-",
                              couleur_hausse=ROUGE if h >= 10 else (JAUNE if h > 0 else BLANC)))
    lignes_zs.sort(key=lambda z: (-z["_cas"], z["nom"]))
    sans_cas = [z["nom"] for z in lignes_zs if z["_cas"] == 0]
    lignes_zs_affichees = [z for z in lignes_zs if z["_cas"] > 0]      # les ZS sans cas sont citées en note
    tot_cgpp = dict(cas=n(tz["cas"]), hausse=n(tz["hausse"], signe=True), nouveaux=n(tz["nouveaux"]),
                    deces=n(tz["deces"]), letalite=pct(100 * tz["deces"] / tz["cas"]) if tz["cas"] else "-")
    post_nat = int(zs1.deces_communautaires_24h.fillna(0).sum())
    nouv_nat = int(zs1.nouveaux_cas_24h.fillna(0).sum())

    # ---------------- 3. données opérationnelles consolidées (dernier Weekly Update)
    ops = res.get("operations", {})
    om, ol = ops.get("meta", {}), {(l["id_indicateur"], l["desagregation"]): l for l in ops.get("lignes", [])}
    cols = om.get("colonnes", [])
    def v(id_, desag, col, role="num"):
        l = ol.get((id_, desag), {})
        x = l.get(f"{role}|{col}")
        return None if x is None or (isinstance(x, float) and pd.isna(x)) else x
    lignes_ops, op = [], {}
    if cols:
        tous = cols + ["Total"]
        def ligne(lib, valeurs, couleurs=None, gras=False):
            lignes_ops.append(dict(cellules=[lib] + valeurs, couleurs=couleurs or {}, gras=gras))
        def ligne_taux(lib, id_, desag, vert, jaune, compte=True):
            vals, coul = [], {}
            for j, c in enumerate(tous, start=1):
                a, b = v(id_, desag, c), v(id_, desag, c, "den")
                t_ = taux(a, b)
                vals.append((f"{n(a)} ({pct(t_)})" if compte else pct(t_)) if t_ is not None else "-")
                coul[j] = couleur_taux(t_, vert, jaune) if t_ is not None else GRIS
            ligne(lib, vals, coul)
        S29, S07, S02, C14 = ("SURV-29", "RECO actifs / RECO attendus"), ("SURV-07", "alertes notifiées / alertes détectées"), \
            ("SURV-02", "contacts vus / contacts suivis (cumul)"), ("CREC-14", "résistances levées / résistances identifiées")
        ligne("RECO actifs / attendus", [f"{n(v(*S29, c))} / {n(v(*S29, c, 'den'))}" for c in tous])
        ligne_taux("Taux de participation des RECO", *S29, 90, 80, compte=False)
        ligne("Rapports SBC reçus", [n(v("SURV-29", "rapports SBC reçus", c)) for c in tous])
        ligne("Personnes sensibilisées", [n(v("CREC-03", "personnes sensibilisées", c)) for c in tous])
        ligne("Alertes (malades) détectées", [n(v("SURV-05", "alertes (malades) détectées", c)) for c in tous])
        ligne_taux("Alertes notifiées (taux)", *S07, 95, 80)
        ligne("Décès communautaires prélevés", [n(v("SURV-30", "décès communautaires prélevés", c)) for c in tous])
        ligne("Enterrements dignes et sécurisés", [n(v("CREC-17", "EDS réalisés", c)) for c in tous])
        ligne("Contacts suivis (cumul)", [n(v("SURV-01", "contacts suivis (cumul)", c)) for c in tous])
        ligne_taux("Contacts vus (taux)", *S02, 90, 80)
        ligne("Résistances identifiées", [n(v(*C14, c, "den")) for c in tous])
        ligne_taux("Résistances levées (taux)", *C14, 80, 70)
        op = dict(reco=taux(v(*S29, "Total"), v(*S29, "Total", "den")), reco_n=v(*S29, "Total"),
                  reco_d=v(*S29, "Total", "den"), sens=v("CREC-03", "personnes sensibilisées", "Total"),
                  al_d=v(*S07, "Total", "den"), al_n=v(*S07, "Total"), al=taux(v(*S07, "Total"), v(*S07, "Total", "den")),
                  ct_n=v(*S02, "Total"), ct_d=v(*S02, "Total", "den"), ct=taux(v(*S02, "Total"), v(*S02, "Total", "den")),
                  res=taux(v(*C14, "Total"), v(*C14, "Total", "den")), res_n=v(*C14, "Total"), res_d=v(*C14, "Total", "den"),
                  prel=v("SURV-30", "décès communautaires prélevés", "Total"), eds=v("CREC-17", "EDS réalisés", "Total"),
                  debut=pd.to_datetime(om["debut"]), fin=pd.to_datetime(om["fin"]))
        faible = [(c, taux(v(*S02, c), v(*S02, c, "den"))) for c in cols]
        op["ct_min"] = min((x for x in faible if x[1] is not None), key=lambda x: x[1], default=(None, None))
        op["ct_min_n"] = (v(*S02, op["ct_min"][0]), v(*S02, op["ct_min"][0], "den")) if op["ct_min"][0] else (None, None)
    entetes_ops = ["Indicateur"] + [c.replace("WV ", "World Vision, ").replace("ADRA ", "ADRA, ") for c in cols] + ["Total CGPP"]

    # ---------------- 4. performance contre les cibles
    ref = lire_table(base, "ref_indicateurs").set_index("id_indicateur")
    saisie = lire_table(base, "saisie_hebdo")
    for c in ("numerateur", "denominateur"):
        saisie[c] = pd.to_numeric(saisie[c], errors="coerce")
    saisie["date_fin"] = pd.to_datetime(saisie["date_fin"])
    comp = pd.read_csv(cfg["chemins"]["intermediaire"] / "completude_classeur.csv")
    def periode_txt(g):
        p_ = g.periode.iloc[0]
        src = "WU" if "Weekly" in str(g.source_saisie.iloc[0]) else "classeur"
        if str(p_).startswith("Cumul"):
            return f"{p_.lower()} ({src})"
        a, b = pd.to_datetime(g.date_debut.iloc[0]), g.date_fin.iloc[0]
        return (f"{a:%d}-{b:%d/%m}" if a.month == b.month else f"{a:%d/%m}-{b:%d/%m}") + f" ({src})"
    indic = []
    LIBELLES = {"SURV-29": "Complétude SBC : RECO actifs / RECO attendus",
                "SURV-07": "Alertes communautaires notifiées / détectées (délai non précisé)",
                "SURV-02": "Contacts vus / contacts suivis (cumul, définition à préciser)",
                "CREC-14": "Résistances communautaires levées (délai non précisé)"}
    def ajoute(id_, valeur, periode, statut, couleur, cible=None):
        r = ref.loc[id_] if id_ in ref.index else None
        lib = LIBELLES.get(id_) or (r.libelle_fr if r is not None else id_)
        indic.append(dict(id=id_, libelle=lib, valeur=valeur, periode=periode,
                          cible=cible if cible is not None else (r.cible_dos if r is not None and pd.notna(r.cible_dos) else "-"),
                          statut=statut, couleur=couleur))
    def derniere(id_, desag=None):
        s = saisie[saisie.id_indicateur == id_]
        if desag:
            s = s[s.desagregation == desag]
        if s.empty:
            return s
        return s[s.date_fin == s.date_fin.max()]
    for id_, desag in [("SURV-29", "RECO actifs / RECO attendus"), ("SURV-07", "alertes notifiées / alertes détectées"),
                       ("SURV-02", "contacts vus / contacts suivis (cumul)"),
                       ("CREC-14", "résistances levées / résistances identifiées")]:
        g = derniere(id_, desag)
        if g.empty:
            ajoute(id_, "-", "-", "Non rapporté", GRIS)
            continue
        val = taux(g.numerateur.sum(), g.denominateur.sum())
        cible = ref.loc[id_].cible_dos_num if pd.notna(ref.loc[id_].cible_dos_num) else 80.0
        statut, coul = (("Cible atteinte", VERT) if val >= cible else
                        (("Proche de la cible", JAUNE) if val >= cible - 10 else ("Sous la cible", ROUGE)))
        ajoute(id_, f"{pct(val)} ({n(g.numerateur.sum())}/{n(g.denominateur.sum())})", periode_txt(g), statut, coul,
               cible=ref.loc[id_].cible_dos if pd.notna(ref.loc[id_].cible_dos) else "≥ 80 % (seuil)")
    ajoute("SURV-22", f"{pct(100 * post_cgpp / nouv_cgpp, 0)} ({int(post_cgpp)}/{int(nouv_cgpp)})" if nouv_cgpp else "-",
           f"{d1:%d/%m} (SitRep)", "À faire baisser", JAUNE if nouv_cgpp and post_cgpp / nouv_cgpp > 0.2 else BLEU_PALE)
    for id_, desag, harmoniser in [("SURV-05", "alertes (malades) détectées", True), ("CREC-03", "personnes sensibilisées", True),
                                   ("SURV-14", None, False), ("SURV-03", None, False), ("SURV-09", None, False),
                                   ("COORD-03", None, False)]:
        g = derniere(id_, desag)
        if g.empty:
            ajoute(id_, "-", "-", "Non rapporté", GRIS)
            continue
        val, cible = g.numerateur.sum(), ref.loc[id_].cible_dos_num
        cumul = str(g.periode.iloc[0]).startswith("Cumul")
        duree = (g.date_fin.iloc[0] - pd.to_datetime(g.date_debut.iloc[0])).days + 1
        if cumul and pd.notna(cible):
            statut, coul = f"{pct(100 * val / cible, 0)} de la cible", BLEU_PALE
        elif harmoniser and pd.notna(cible) and duree <= 7 and val >= 0.25 * cible:
            statut, coul = f"{pct(100 * val / cible, 0)} de la cible en {duree} j : définition à revoir", JAUNE
        else:
            statut, coul = "Cumul à consolider", JAUNE
        ajoute(id_, n(val), periode_txt(g), statut, coul)
    v_comp = 100 * comp.renseigne.mean()
    ajoute("MEL-15", f"{pct(v_comp)} ({int(comp.renseigne.sum())}/{len(comp)})", f"S36-S{int(comp.semaine.max())} (classeur)",
           "Sous le seuil de 80 %" if v_comp < 80 else "Conforme", ROUGE if v_comp < 80 else VERT, cible="≥ 80 % (seuil)")

    # ---------------- notes, autres épidémies, points d'action, qualité
    notes = lire_table(base, "sitrep_notes")
    notes["date_rapport"] = pd.to_datetime(notes["date_rapport"])
    dn = notes[notes.date_rapport <= d1].date_rapport.max()
    nn = notes[notes.date_rapport == dn].sort_values("ordre")
    def notes_de(section, defaut="Aucune note saisie pour cette date dans la feuille sitrep_notes."):
        v_ = list(nn[nn.section == section]["texte"])
        return v_ or [defaut]
    ae = lire_table(base, "autres_epidemies")
    epi = []
    if len(ae):
        ae = ae[ae.semaine_epi == ae.semaine_epi.max()]
        for r in ae.sort_values(["maladie", "province"]).itertuples():
            epi.append(dict(maladie=r.maladie, province=r.province,
                            cas=f"{n(r.cas)} ({n(r.deces) if pd.notna(r.deces) else 'ND'})",
                            zs=f"{int(r.zs_notifiantes)}/{int(r.zs_total)}" if pd.notna(r.zs_notifiantes) else "-",
                            cgpp=r.zs_cgpp_concernees if pd.notna(r.zs_cgpp_concernees) else "-",
                            ref=f"{n(r.cas_suspects_referes_cgpp)} ({r.detail_references})"
                            if pd.notna(r.cas_suspects_referes_cgpp) else "-"))
    semaine_epi = int(ae.semaine_epi.max()) if len(ae) else None
    pa = lire_table(base, "points_action").merge(lire_table(base, "reunions")[["id_reunion", "date", "niveau"]],
                                                    on="id_reunion", how="left")
    coul_action = {"À faire": JAUNE, "En cours": BLEU_PALE, "Réalisé": VERT, "Abandonné": GRIS}
    actions = [dict(reunion=f"{pd.to_datetime(r.date):%d/%m} ({r.niveau})" if pd.notna(r.date) else r.id_reunion,
                    action=r.action, responsable=r.responsable if pd.notna(r.responsable) else "À désigner",
                    echeance=f"{pd.to_datetime(r.echeance):%d/%m/%Y}" if pd.notna(r.echeance) else "À fixer",
                    statut=r.statut, couleur=coul_action.get(r.statut, GRIS)) for r in pa.itertuples()]
    qc_fichiers = sorted(cfg["chemins"]["sorties_qc"].glob("rapport_qc_*.xlsx"))
    texte_qc, anomalies = "Aucun rapport de contrôle qualité trouvé.", []
    if qc_fichiers:
        tous_c = pd.read_excel(qc_fichiers[-1], sheet_name="tous_les_controles")
        an = pd.read_excel(qc_fichiers[-1], sheet_name="anomalies")
        g = an.groupby("gravite").size()
        texte_qc = (f"{len(tous_c)} contrôles automatiques : {(tous_c.statut == 'OK').sum()} conformes, "
                    f"{(tous_c.statut == 'ÉCHEC').sum()} échecs et {(tous_c.statut == 'À VÉRIFIER').sum()} points à "
                    f"vérifier ({g.get('Critique', 0)} critiques, {g.get('Majeure', 0)} majeures, "
                    f"{g.get('Mineure', 0)} mineures). À corriger en priorité :")
        def court(x, m=105):
            x = str(x)
            return x if len(x) <= m else x[:m].rsplit(" ", 1)[0] + " (…)"
        sel = an[an.gravite.isin(["Critique", "Majeure"])].copy()
        sel["_g"] = sel.gravite.map({"Critique": 0, "Majeure": 1})
        priorite = ["Weekly Updates CGPP", "Classeur hebdomadaire et cadre global", "SitReps nationaux",
                    "Plans de travail (DOS, Tranche 5)", "Comptes rendus de coordination"]
        sel["_f"] = sel.famille.map({f: i for i, f in enumerate(priorite)}).fillna(9)
        premieres = sel.sort_values(["_g", "_f"]).drop_duplicates("id_regle")
        choix = list(premieres[premieres._f < 4].drop_duplicates("famille").sort_values(["_f", "_g"]).id_regle)[:4]
        for rid in choix:
            gg = sel[sel.id_regle == rid]
            r = gg.iloc[0]
            autres = f" ; {len(gg) - 1} autre(s) cas" if len(gg) > 1 else ""
            anomalies.append(f"[{rid}, {r.gravite.lower()}] {r.source} : {r.objet} ; attendu {court(r.attendu, 50)}, "
                             f"observé {court(r.observe)}{autres}.")
        reste = [x for x in premieres.id_regle if x not in choix]
        if reste:
            anomalies.append(f"{len(reste)} autres règles critiques ou majeures en échec : détail dans "
                             f"{qc_fichiers[-1].name} (feuille anomalies).")

    # ---------------- textes rédigés par règles
    t_zs = pd.read_excel(cfg["chemins"]["sorties_analyses"] / "analyses.xlsx", sheet_name="zones")
    t_prov = pd.read_excel(cfg["chemins"]["sorties_analyses"] / "analyses.xlsx", sheet_name="provinces").set_index("province")
    foyers = t_zs[t_zs.foyer_actif.astype(bool)].sort_values("hausse", ascending=False)
    def foyer_txt(r):
        return (f"{r.zone_sante} ({n(r.hausse, signe=True)} cas" + (f", {n(r.hausse_pct, 0, signe=True)} %)" if pd.notna(r.hausse_pct) else ")"))
    hors = foyers[foyers.ip_cgpp.isna()]
    nouvelles_zs = sorted(set(zs1[~zs1.est_a_ventiler.fillna(False).astype(bool)].zone_sante) -
                          set(zs0[~zs0.est_a_ventiler.fillna(False).astype(bool)].zone_sante))
    recents = [i for i in ivs if pd.to_datetime(i["debut"]) >= (d_ref if d_ref is not None else d1)]
    rythmes = [i["jour_cas_National"] for i in recents] or [nat["moyenne_jour"]]
    long_iv = max(ivs, key=lambda i: i["jours"]) if ivs else None
    part_nk_q = 100 * q["provinces"].get("Nord-Kivu", 0) / q["cas"] if q and q.get("cas") else None
    part_nk_long = 100 * long_iv["cas_Nord-Kivu"] / long_iv["cas_National"] if long_iv else None
    nk_recents = [i["jour_cas_Nord-Kivu"] for i in recents]
    tshopo = (q["provinces"].get("Tshopo"), q["cas_ref"].get("Tshopo")) if q else (None, None)

    messages = []
    if q:
        moy_q, moy_l = q["cas"] / q["jours"], (long_iv["jour_cas_National"] if long_iv else None)
        compar = ""
        if moy_l:
            ecart = 100 * (moy_q - moy_l) / moy_l
            compar = (f"C'est un peu moins que la moyenne du {jour_mois(long_iv['debut'])} au {jour_mois(long_iv['fin'])} "
                      f"({n(moy_l, 0)} par jour)" if ecart < -5 else
                      (f"C'est plus que la moyenne du {jour_mois(long_iv['debut'])} au {jour_mois(long_iv['fin'])} "
                       f"({n(moy_l, 0)} par jour)" if ecart > 5 else
                       f"C'est le niveau moyen du {jour_mois(long_iv['debut'])} au {jour_mois(long_iv['fin'])}"))
        baisse_continue = len(rythmes) > 1 and all(b < a for a, b in zip(rythmes, rythmes[1:]))
        suite = ", puis ".join(n(x, 0) for x in rythmes)
        messages.append(
            f"L'épidémie reste très active : {n(q['cas'])} nouveaux cas confirmés et {n(q['deces'])} décès en "
            f"{q['jours']} jours (du {jour_mois(d_ref)} au {jour_mois(d1)}), soit {n(moy_q, 0)} cas par jour. " +
            (compar + (", et la baisse est continue" if baisse_continue else ", mais sans baisse continue") +
             f" : {suite} cas par jour sur les {len(rythmes)} derniers intervalles." if compar else ""))
        messages.append(
            f"Le Nord-Kivu concentre {pct(part_nk_q, 0)} des nouveaux cas de la période "
            f"({n(q['provinces'].get('Nord-Kivu', 0))} sur {n(q['cas'])}), contre {pct(part_nk_long, 0)} du "
            f"{jour_mois(long_iv['debut'])} au {jour_mois(long_iv['fin'])} ; il ajoute désormais de {n(min(nk_recents), 0)} "
            f"à {n(max(nk_recents), 0)} cas par jour. "
            f"L'Ituri reste l'épicentre ({n(t_prov.loc['Ituri', 'cas'])} cas, {pct(t_prov.loc['Ituri', 'part_cas_pct'], 0)} "
            f"du cumul).")
    tx_tshopo = (f" ; la Tshopo passe de {n(tshopo[1])} à {n(tshopo[1] + tshopo[0])} cas en {q['jours']} jours "
                 f"({n(100 * tshopo[0] / tshopo[1], 0, signe=True)} %)" if tshopo[0] and tshopo[1] else "")
    messages.append(f"Foyers actifs sur les {jours} derniers jours hors du dispositif CGPP : "
                    f"{', '.join(foyer_txt(r) for r in hors.itertuples())}{tx_tshopo}.")
    messages.append(f"La détection reste tardive : {post_nat} des {nouv_nat} nouveaux cas du {jour_mois(d1)} "
                    f"({pct(100 * post_nat / nouv_nat, 0) if nouv_nat else 'ND'}) ont été confirmés après le décès ; dans "
                    f"les ZS du dispositif CGPP, c'est le cas de {int(post_cgpp)} des {int(nouv_cgpp)} nouveaux cas "
                    f"({pct(100 * post_cgpp / nouv_cgpp, 0) if nouv_cgpp else 'ND'}).")
    if op:
        min_c = op["ct_min"]
        messages.append(
            f"Du {op['debut'].day} au {jour_mois(op['fin'])}, {n(op['reco_n'])} RECO actifs sur {n(op['reco_d'])} "
            f"({pct(op['reco'])}) ont sensibilisé {n(op['sens'])} personnes et détecté {n(op['al_d'])} alertes, dont "
            f"{pct(op['al'], 0)} notifiées ; {n(op['prel'])} décès communautaires ont été prélevés et {n(op['eds'])} EDS "
            f"réalisés. Le suivi des contacts reste insuffisant en {min_c[0].split(' ', 1)[1]} ({min_c[0].split(' ')[0]} : "
            f"{n(op['ct_min_n'][0])} contacts vus sur {n(op['ct_min_n'][1])}, {pct(min_c[1], 0)})." if min_c[0] else "")
    secu = notes_de("Sécurité", "")
    if secu and secu[0]:
        t0 = secu[0]
        messages.append("Sécurité : " + (t0[0].lower() + t0[1:] if len(t0) > 1 and t0[1].islower() else t0))
    messages = [m for m in messages if m]

    jours_retard = (d1 - op["fin"]).days if op else None
    presents = {c.split(" ")[0] for c in cols}
    absents = sorted(set(cfg.get("partenaires", {})) - presents)
    completude = (f"SitRep national n° {no1} du {d1:%d/%m} : {len(c1) - len(nd_prov)} provinces sur {len(c1)} ont "
                  f"transmis alertes et suivi des contacts" + (f" ({', '.join(nd_prov)} : non transmis)" if nd_prov else "") +
                  ". Données opérationnelles CGPP : " +
                  (f"Weekly Update du {op['fin']:%d/%m} ({', '.join(cols)})" +
                   (f" ; {liste_fr(absents)} absents" if absents else "") +
                   f" ; données vieilles de {jours_retard} jours au moment du SitRep. " if op else "aucune. ") +
                  f"Classeur hebdomadaire par ZS : {res['cellules_renseignees']} cellules renseignées sur "
                  f"{res['cellules_attendues']} ({pct(res['completude_classeur'])}).")
    texte_epi = (f"Au {jour_mois(d1, True)} (SitRep n° {no1}), la RDC cumule {n(nat['cas'])} cas confirmés et "
                 f"{n(nat['deces'])} décès (létalité {pct(nat['letalite'])}), {n(nat['gueris'])} guéris et "
                 f"{n(nat['isolement'])} patients en isolement, dans {nat['zs_touchees']} zones de santé sur "
                 f"{nat['zs_total']} et {nat['provinces']} provinces. La figure 1 compare le rythme des nouveaux cas "
                 f"entre les {len(ivs) + 1} SitReps disponibles, dont les n° " +
                 liste_fr([int(i["sitrep_fin"]) for i in ivs if int(i["sitrep_fin"]) not in (no0, no1)]) +
                 " cités par les Weekly Updates CGPP : le rythme reste élevé, sans baisse continue, et la part du Nord-Kivu "
                 "augmente." +
                 (f" Nouvelle zone de santé touchée le {jour_mois(d1)} : {', '.join(nouvelles_zs)}." if nouvelles_zs else ""))
    hausses_cgpp = [z for z in sorted(lignes_zs, key=lambda z: -z["_h"]) if z["_h"] > 0][:3]
    texte_cgpp = (f"Les {cg['zs_total']} ZS du dispositif comptent {n(cg['cas'])} cas confirmés et {n(cg['deces'])} décès "
                  f"({pct(cg['part_cas_pct'])} du cumul national) et {n(cg['hausse'])} des {n(nat['nouveaux_periode'])} "
                  f"cas des {jours} derniers jours ({pct(cg['part_hausse_pct'])}). " +
                  liste_fr([f"{z['nom']} ({z['hausse']})" for z in hausses_cgpp]) +
                  f" portent l'essentiel de cette hausse. {len(cg['zs_sans_cas'].split(', '))} ZS "
                  f"({liste_fr(cg['zs_sans_cas'].split(', '))}) n'ont aucun cas et relèvent de la préparation.")
    ecarts = ecarts_publies(cfg, op["fin"]) if op else []
    texte_ops = ((f"Tableau consolidé du Weekly Update du {op['fin']:%d/%m} (activités du {op['debut']:%d} au "
                  f"{op['fin']:%d/%m}), seule source opérationnelle disponible pour la période. Les totaux et les taux "
                  f"sont recalculés à partir des colonnes des partenaires" +
                  (f" : le rapport publiait {liste_fr(ecarts)}." if ecarts else ".")) if op
                 else "Aucun tableau opérationnel disponible.")
    tuiles = dict(
        t1=dict(label="Cas confirmés cumulés", valeur=n(nat["cas"]),
                sous=(f"{n(q['cas'], signe=True)} en {q['jours']} j ({n(q['cas'] / q['jours'], 0)}/j)" if q else "")),
        t2=dict(label="Décès confirmés", valeur=n(nat["deces"]),
                sous=f"létalité {pct(nat['letalite'])} ; {n(q['deces'], signe=True)} en {q['jours']} j" if q else ""),
        t3=dict(label=f"Cas par jour (du {(d1 - pd.Timedelta(days=jours)).day} au {d1:%d/%m})",
                valeur=n(nat["moyenne_jour"], 1),
                sous=(f"{n(ivs[-2]['jour_cas_National'], 1)} du {pd.to_datetime(ivs[-2]['debut']):%d/%m} au "
                      f"{pd.to_datetime(ivs[-2]['fin']):%d/%m}" if len(ivs) > 1 else "")),
        t4=dict(label="Zones de santé touchées", valeur=f"{nat['zs_touchees']} / {nat['zs_total']}",
                sous=(f"{n(nat['zs_touchees'] - q['zs_ref'], signe=True)} en {q['jours']} j" if q else "") +
                     (f" ; dernière : {', '.join(nouvelles_zs)}" if nouvelles_zs else "")),
        t5=dict(label="Cas dans les 12 ZS CGPP", valeur=n(cg["cas"]),
                sous=f"{n(cg['hausse'], signe=True)} en {jours} j ({pct(cg['part_hausse_pct'], 0)} de la hausse)"),
        t6=dict(label=f"RECO actifs ({op['debut']:%d}-{op['fin']:%d/%m})" if op else "RECO actifs",
                valeur=pct(op.get("reco")), sous=f"{n(op.get('reco_n'))} sur {n(op.get('reco_d'))} attendus"),
        t7=dict(label="Alertes communautaires notifiées", valeur=pct(op.get("al")),
                sous=f"{n(op.get('al_n'))} sur {n(op.get('al_d'))} détectées"),
        t8=dict(label="Contacts vus / suivis (CGPP)", valeur=pct(op.get("ct")),
                sous=(f"{op['ct_min'][0].split(' ')[0]} {op['ct_min'][0].split(' ', 1)[1]} : {pct(op['ct_min'][1], 0)}"
                      if op and op["ct_min"][0] else "")),
    )
    semaine = int(d1.isocalendar().week)
    sources = (f"SitReps nationaux MVE n° {no0} ({m0.date_rapport:%d/%m/%Y}) et {no1} ({d1:%d/%m/%Y}), INSP/SGI MVE-17 ; "
               f"SitReps n° " + liste_fr([int(i["sitrep_fin"]) for i in ivs if int(i["sitrep_fin"]) not in (no0, no1)]) +
               " cités par les Weekly Updates CGPP des 28/09 et 30/09/2026 ; SitRep n° 085 (7 août) cité par le plan "
               "de travail DOS ; classeur hebdomadaire CGPP ; comptes rendus de coordination nationale (05/10) et de "
               "l'Ituri (06/10) ; contrôle qualité automatisé.")
    sigles = ("ACSA : agent de santé communautaire ; CTE : centre de traitement Ebola ; EDS : enterrement digne et "
              "sécurisé ; EIR : équipe d'intervention rapide ; IP : partenaire de mise en oeuvre ; PoC/PoE : point de "
              "contrôle, point d'entrée ; RECO : relais communautaire ; SBC : surveillance à base communautaire ; "
              "SONUB : soins obstétricaux et néonatals d'urgence de base ; WU : Weekly Update CGPP ; ZS : zone de santé.")
    return dict(
        **tuiles,
        numero=f"S{semaine}-{d1.year}", date_donnees=f"{d1:%d/%m/%Y}", sitrep_national=no1,
        periode_ops=(f"{op['debut']:%d}-{op['fin']:%d/%m/%Y}" if op else "ND"),
        date_publication=f"{datetime.now():%d/%m/%Y}", date_generation=f"{datetime.now():%d/%m/%Y à %H:%M}",
        jours=jours, j14=q.get("jours", "-"), date_ref=f"{d_ref:%d/%m}" if d_ref is not None else "-",
        completude_texte=completude, messages=messages, texte_epi=texte_epi, provinces=lignes_prov, tot=tot,
        note_provinces=(f"Hausse {q.get('jours', '')} j : depuis le SitRep n° {q.get('sitrep_ref', '')} du "
                        f"{d_ref:%d/%m} (cité par le Weekly Update CGPP) ; hausse {jours} j : depuis le SitRep "
                        f"n° {no0}. Contacts vus : vert ≥ 95 %, jaune de 90 à 95 %, rouge < 90 %, ND : non transmis. "
                        "En Ituri, 493 décès ne sont pas encore attribués à une zone de santé.") if q else "",
        legende_figure1=("Nouveaux cas confirmés par jour, en moyenne sur chaque intervalle entre deux SitReps nationaux "
                         "(n° 085 cité par le plan DOS ; n° 135 et 138 cités par les Weekly Updates)."),
        texte_cgpp=texte_cgpp, zs_cgpp=lignes_zs_affichees, tot_cgpp=tot_cgpp,
        note_cgpp=(f"Hausse sur {jours} jours : rouge ≥ 10 cas, jaune 1 à 9 cas. Total : les {len(lignes_zs)} ZS, dont "
                   f"{len(sans_cas)} sans cas non affichées ({liste_fr(sans_cas)}). Foyers actifs hors dispositif "
                   f"(hausse ≥ 5 cas et ≥ 15 %, ou ≥ 10 cas) : {cv['foyers_hors_cgpp']}."),
        texte_ops=texte_ops, entetes_ops=entetes_ops, lignes_ops=lignes_ops,
        note_ops=("Couleurs : vert, jaune, rouge selon le seuil de chaque taux (participation 90 et 80 %, notification "
                  "95 et 80 %, contacts 90 et 80 %, résistances 80 et 70 %). Contacts : le rapport donne un cumul de "
                  "contacts suivis et de contacts vus, pas le suivi du jour ; le taux n'est donc pas comparable au "
                  "suivi quotidien du SitRep national."),
        notes_faits=notes_de("Faits marquants", ""), notes_activites=notes_de("Activités"),
        texte_performance=("Dernière valeur disponible de chaque indicateur prioritaire, avec sa période et sa source. Les "
                           "taux se comparent à la cible ; les nombres d'une seule période ne se comparent pas à une "
                           "cible fixée pour toute la durée du plan."),
        indicateurs=indic,
        notes_secu=notes_de("Sécurité"), notes_defis=notes_de("Défis"), notes_actions=notes_de("Actions"),
        texte_epidemies=(f"Situation de la semaine épidémiologique {semaine_epi} dans les provinces d'intervention et cas "
                         "suspects alertés et référés par les RECO appuyés par CGPP." if semaine_epi else ""),
        epidemies=epi,
        note_epidemies=("Cas (décès) : semaine épidémiologique ; ZS : ZS ayant notifié des cas sur l'ensemble de la "
                        "province. Les deux Weekly Updates reprennent la même section sans mise à jour."),
        actions=actions, texte_qc=texte_qc, anomalies=anomalies or ["Aucune anomalie critique ou majeure."],
        sources=sources, sigles=sigles,
    ), d1


def main(cfg: dict | None = None, date_cible: str | None = None):
    cfg = cfg or charger_config()
    ctx, d1 = contexte(cfg, date_cible)
    tpl = DocxTemplate(cfg["chemins"]["canevas_sitrep"])
    fig = cfg["chemins"]["sorties_figures"] / "f6_rythme.png"
    ctx["figure_rythme"] = InlineImage(tpl, str(fig), width=Mm(150)) if fig.exists() else ""
    ctx["marque_operations"] = PLACE_OPERATIONS
    tpl.render(ctx, autoescape=True)
    sortie = cfg["chemins"]["sorties_sitrep"] / f"SitRep_CGPP_{d1:%Y-%m-%d}.docx"
    tpl.save(sortie)
    # tableau à colonnes variables (une colonne par IP et province présents dans le dernier Weekly Update)
    doc = docx.Document(sortie)
    if ctx["lignes_ops"]:
        k = len(ctx["entetes_ops"]) - 1
        largeurs = [52] + [round(118 / k, 1)] * k
        inserer_tableau(doc, PLACE_OPERATIONS, ctx["entetes_ops"], ctx["lignes_ops"], largeurs)
    else:
        for p in doc.paragraphs:
            if p.text.strip() == PLACE_OPERATIONS:
                p.text = "Aucune donnée opérationnelle consolidée disponible."
    ordonner_document(doc)
    doc.save(sortie)
    pdf = vers_pdf(sortie, cfg["chemins"]["sorties_sitrep"], cfg.get("libreoffice", "soffice"))
    LOG.info("SitRep CGPP généré : %s et %s", sortie.name, pdf.name)
    return sortie, pdf


if __name__ == "__main__":
    main(date_cible=sys.argv[1] if len(sys.argv) > 1 else None)
